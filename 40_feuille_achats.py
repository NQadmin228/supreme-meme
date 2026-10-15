#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TABOO — Feuille de saisie des achats, article par article.

  python 40_feuille_achats.py
  python 40_feuille_achats.py --nomenclature chemin.xlsx --sortie chemin.xlsx

POURQUOI CETTE FEUILLE
----------------------
Le livre de depenses est tenu par POSTE : une colonne EAT, une DRINK,
une SMOKE, une GAZ, une GLACONS. L'huile, le gigot, le riz, les
emballages, les fruits, les sirops, le charbon ou les aromes y sont
fondus. On ne peut donc pas dire combien on en achete, a quel rythme,
ni si c'est proportionne aux ventes -- les trois questions posees par
la direction. La page « Achats matiere » du tableau de bord y repond au
grain du poste ; cette feuille fournit le grain article.

CE QU'ELLE CONTIENT
-------------------
    Lisez-moi            mode d'emploi, legende, une ligne d'exemple
    Saisie               une ligne par achat ; colonnes jaunes a remplir,
                         colonnes grises calculees
    Articles             le catalogue, tire de la nomenclature des achats
    Synthese prioritaires  montant mois par mois, nombre d'achats et
                         intervalle moyen des articles suivis en priorite
    Controle livre       total saisi par poste et par mois, face au livre
                         de depenses : un ecart dit qu'un achat manque

Le catalogue vient de la nomenclature tenue par l'etablissement : ce
script n'y ajoute que ce qui manque pour qu'une journee se saisisse en
entier (glacons, et une ligne « Autre achat » par activite).

Le fichier produit n'est pas versionne : il porte les prix d'achat
negocies avec les fournisseurs, et ensuite les achats eux-memes.
"""

import argparse
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = Path(__file__).resolve().parent
SORTIE = BASE / "TABOO_Saisie_achats.xlsx"

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# Lignes de saisie preparees. A raison d'une dizaine d'achats par jour
# d'ouverture, c'est environ un an ; au-dela, on recopie la derniere
# ligne vers le bas et les formules suivent.
LIGNES = 3000

# Postes du livre de depenses auxquels un achat peut se rattacher. Ce
# sont les noms de colonnes du livre, a l'identique : c'est ce qui
# permet le controle mensuel.
POSTES = ["EAT", "GAZ", "DRINK", "GLACONS", "SMOKE", "CONSOMMABLES"]

# Les articles nommes dans la demande de la direction (zoom achats).
# Les fruits et sirops du bar, le charbon et les aromes y entrent par
# leur categorie entiere.
PRIORITAIRES = {"Huile", "Gaz (bouteille)", "Gigot", "Riz", "Emballages alu",
                "Emballages papier", "Sucre à café", "Thé", "Boîtes de thon"}
CATEGORIES_PRIORITAIRES = ("DRINK - Fruits", "DRINK - Sirops",
                           "SMOKE - Charbon", "SMOKE - Arômes")

# Ce que la nomenclature n'a pas et qu'une journee d'achats contient.
AJOUTS = [
    ("DRINK - Glaçons", "Glaçons", None, None),
    ("EAT - Divers", "Autre achat EAT (préciser en remarque)", None, None),
    ("DRINK - Divers", "Autre achat DRINK (préciser en remarque)", None, None),
    ("SMOKE - Divers", "Autre achat SMOKE (préciser en remarque)", None, None),
]

PAIEMENTS = "Espèces,Flooz,Mixx,Carte,Virement,Crédit fournisseur"

# ---------------------------------------------------------------------
# Mise en forme
# ---------------------------------------------------------------------
POLICE = "Arial"
F_NORMAL = Font(name=POLICE, size=10)
F_GRAS = Font(name=POLICE, size=10, bold=True)
F_TITRE = Font(name=POLICE, size=14, bold=True, color="12211F")
F_ENTETE = Font(name=POLICE, size=10, bold=True, color="FFFFFF")
F_SAISIE = Font(name=POLICE, size=10, color="0000FF")   # valeur tapee
F_CALCUL = Font(name=POLICE, size=10, color="000000")   # formule
F_LIEN = Font(name=POLICE, size=10, color="008000")     # formule vers une autre feuille
F_NOTE = Font(name=POLICE, size=9, italic=True, color="595959")

FOND_ENTETE = PatternFill("solid", fgColor="12211F")
FOND_SAISIE = PatternFill("solid", fgColor="FFF2CC")
FOND_CALCUL = PatternFill("solid", fgColor="F2F2F2")
FOND_PRIO = PatternFill("solid", fgColor="E2EFDA")
TRAIT = Side(style="thin", color="BFBFBF")
CADRE = Border(left=TRAIT, right=TRAIT, top=TRAIT, bottom=TRAIT)
ENROULE = Alignment(wrap_text=True, vertical="top")

FMT_F = '#,##0 "F";-#,##0 "F";"-"'
FMT_Q = "General"     # « #,##0.## » affiche « 2. » pour un entier
FMT_N = '#,##0'
FMT_D = "dd/mm/yyyy"
FMT_MOIS = "mmm yyyy"
FMT_PCT = "0%"


def entete(ws, ligne, titres, largeurs):
    for i, (t, l) in enumerate(zip(titres, largeurs), start=1):
        c = ws.cell(row=ligne, column=i, value=t)
        c.font, c.fill, c.border = F_ENTETE, FOND_ENTETE, CADRE
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = l
    ws.row_dimensions[ligne].height = 30


def prix(valeur):
    """'1000 a 1500' -> (1000, 1500) ; 2000 -> (2000, 2000) ; vide -> (None, None)."""
    if valeur is None or str(valeur).strip() == "":
        return None, None
    if isinstance(valeur, (int, float)):
        return float(valeur), float(valeur)
    nombres = [float(n) for n in re.findall(r"\d+(?:[.,]\d+)?", str(valeur).replace(",", "."))]
    if not nombres:
        return None, None
    return min(nombres), max(nombres)


def poste_du_livre(categorie, article):
    if article.startswith("Gaz"):
        return "GAZ"
    if article == "Glaçons":
        return "GLACONS"
    return categorie.split(" - ")[0].strip()


def lire_nomenclature(chemin):
    ws = load_workbook(chemin, data_only=True)["Nomenclature achats"]
    articles, vus = [], set()
    for cat, art, _unite_ref, dernier, quantite, *_ in ws.iter_rows(min_row=2, values_only=True):
        if not art or not cat:
            continue
        art, cat = str(art).strip(), str(cat).strip()
        if art.lower() in vus:
            continue
        vus.add(art.lower())
        pmin, pmax = prix(dernier)
        articles.append({"art": art, "cat": cat, "poste": poste_du_livre(cat, art),
                         "unite": (str(quantite).strip() if quantite else None),
                         "min": pmin, "max": pmax, "source": "Nomenclature achats"})
    for cat, art, pmin, pmax in AJOUTS:
        if art.lower() not in vus:
            articles.append({"art": art, "cat": cat, "poste": poste_du_livre(cat, art),
                             "unite": None, "min": pmin, "max": pmax,
                             "source": "Ajouté pour la saisie"})
    for a in articles:
        a["prio"] = a["art"] in PRIORITAIRES or a["cat"].startswith(CATEGORIES_PRIORITAIRES)
    # Ordre de la liste deroulante : par poste, puis categorie, puis nom.
    rang = {p: i for i, p in enumerate(POSTES)}
    articles.sort(key=lambda a: (rang.get(a["poste"], 99), a["cat"], a["art"].lower()))
    return articles


# ---------------------------------------------------------------------
# Feuilles
# ---------------------------------------------------------------------
def feuille_lisez_moi(ws, articles):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    for col, l in zip("BCDEFGH", (14, 30, 12, 14, 18, 14, 34)):
        ws.column_dimensions[col].width = l
    ws["B2"] = "TABOO — Saisie des achats, article par article"
    ws["B2"].font = F_TITRE
    lignes = [
        ("Pourquoi", "Le livre de dépenses ne distingue que les postes (EAT, DRINK, SMOKE, GAZ, "
                     "GLACONS). Pour savoir combien on achète d'huile, de gigot ou de charbon, à "
                     "quel rythme, et si c'est proportionné aux ventes, il faut noter chaque achat "
                     "article par article. C'est le rôle de ce classeur."),
        ("Qui, quand", "La personne qui fait les courses, le jour même, à partir du ticket ou de "
                       "la facture. Une ligne par article acheté : un ticket de 6 articles donne "
                       "6 lignes."),
        ("Où écrire", "Uniquement dans les cellules JAUNES de la feuille « Saisie » (colonnes A à "
                      "G), de « Contrôle livre » (montants du livre) et dans la colonne "
                      "« Poste du livre » de « Articles ». Tout le reste se calcule."),
        ("Article", "Choisir dans la liste déroulante. Un article absent de la liste : prendre "
                    "« Autre achat … » de la bonne activité et le nommer en remarque, puis "
                    "l'ajouter à la feuille « Articles » à la fin du mois."),
        ("Quantité", "En UNITÉS D'ACHAT, telles qu'écrites dans la colonne « Unité d'achat » : "
                     "deux bouteilles de gaz de 25 kg → 2 ; un demi-sac de riz de 10 kg → 0,5."),
        ("Montant", "Ce qui a été réellement payé, en francs CFA, pour la ligne entière "
                    "(pas le prix unitaire)."),
        ("Contrôle prix", "Compare le prix payé par unité à la fourchette de la nomenclature, "
                          "avec la tolérance réglée dans « Articles ». « Au-dessus » ou « en "
                          "dessous » ne veut pas dire erreur : soit le prix a changé, soit la "
                          "quantité ou l'unité a été mal notée."),
        ("Contrôle livre", "Chaque fin de mois, reporter le total de chaque poste du livre de "
                           "dépenses : l'écart montre ce qui n'a pas été saisi ici."),
    ]
    r = 4
    for titre, texte in lignes:
        ws.cell(row=r, column=2, value=titre).font = F_GRAS
        c = ws.cell(row=r, column=3, value=texte)
        c.font, c.alignment = F_NORMAL, ENROULE
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
        ws.row_dimensions[r].height = 15 * max(1, len(texte) // 95 + 1)
        r += 1

    r += 1
    ws.cell(row=r, column=2, value="Légende").font = F_GRAS
    r += 1
    for fond, police, texte in ((FOND_SAISIE, F_SAISIE, "à remplir"),
                                (FOND_CALCUL, F_CALCUL, "calculé, ne pas modifier"),
                                (FOND_PRIO, F_NORMAL, "article suivi en priorité")):
        c = ws.cell(row=r, column=2, value="Exemple")
        c.fill, c.font, c.border = fond, police, CADRE
        ws.cell(row=r, column=3, value=texte).font = F_NORMAL
        r += 1

    r += 1
    ws.cell(row=r, column=2, value="Une ligne d'exemple (ne pas recopier dans « Saisie »)").font = F_GRAS
    r += 1
    titres = ["Date", "Article", "Quantité", "Montant payé", "Fournisseur", "Paiement", "Remarque"]
    # Le montant de l'exemple est lu dans la nomenclature, pas ecrit ici :
    # ce fichier est versionne dans un depot public, les prix d'achat non.
    gaz = next((a for a in articles if a["poste"] == "GAZ" and a["min"]), None)
    valeurs = ["14/10/2026", gaz["art"] if gaz else "Gaz (bouteille)", 2,
               2 * gaz["min"] if gaz else None, "Dépôt de gaz du quartier",
               "Espèces", "2 bouteilles"]
    for i, (t, v) in enumerate(zip(titres, valeurs), start=2):
        h = ws.cell(row=r, column=i, value=t)
        h.font, h.fill, h.border = F_ENTETE, FOND_ENTETE, CADRE
        c = ws.cell(row=r + 1, column=i, value=v)
        c.font, c.fill, c.border = F_SAISIE, FOND_SAISIE, CADRE
        if i == 5:
            c.number_format = FMT_F
    r += 3
    note = ws.cell(row=r, column=2,
                   value="Généré par 40_feuille_achats.py à partir de la nomenclature des "
                         "achats. Pour mettre le catalogue à jour, modifier la nomenclature "
                         "puis relancer le script — ou ajouter les lignes directement dans "
                         "« Articles ».")
    note.font = F_NOTE


def feuille_articles(ws, articles):
    titres = ["Article", "Poste du livre", "Catégorie", "Unité d'achat",
              "Prix de référence min (F)", "Prix de référence max (F)",
              "Suivi prioritaire", "Source"]
    entete(ws, 1, titres, [38, 15, 26, 22, 14, 14, 11, 22])
    for r, a in enumerate(articles, start=2):
        vals = [a["art"], a["poste"], a["cat"], a["unite"], a["min"], a["max"],
                "Oui" if a["prio"] else "", a["source"]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = F_SAISIE if c in (2, 5, 6) else F_NORMAL, CADRE
            if c == 2:
                cell.fill = FOND_SAISIE
            if c in (5, 6):
                cell.number_format = FMT_F
            if a["prio"] and c in (1, 7):
                cell.fill = FOND_PRIO
    der = len(articles) + 1
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:H{der}"
    dv = DataValidation(type="list", formula1='"' + ",".join(POSTES) + '"', allow_blank=False)
    dv.error, dv.errorTitle = "Choisir un poste du livre de dépenses.", "Poste inconnu"
    ws.add_data_validation(dv)
    dv.add(f"B2:B{der + 200}")
    ws["B1"].comment = Comment(
        "Colonne du livre de dépenses où cet achat est noté. Déduit de la catégorie de "
        "la nomenclature : à confirmer avec l'établissement, en particulier pour les "
        "emballages (EAT ou CONSOMMABLES ?).", "40_feuille_achats.py")
    ws["E1"].comment = Comment(
        "Dernier prix connu de la nomenclature, pour une unité d'achat. Une fourchette "
        "« 1000 a 1500 » donne min 1000 et max 1500.", "40_feuille_achats.py")

    # Tolerance du controle de prix : une hypothese, donc une cellule.
    ws["J1"], ws["K1"] = "Tolérance du contrôle de prix", 0.10
    ws["J1"].font, ws["K1"].font = F_GRAS, F_SAISIE
    ws["K1"].fill, ws["K1"].number_format, ws["K1"].border = FOND_SAISIE, FMT_PCT, CADRE
    ws["J2"] = ("Un prix payé jusqu'à 10 % hors de la fourchette reste « dans la fourchette » : "
                "les prix du marché bougent d'un achat à l'autre.")
    ws["J2"].font, ws["J2"].alignment = F_NOTE, ENROULE
    ws.column_dimensions["J"].width = 30
    ws.row_dimensions[2].height = 52
    return der


def feuille_saisie(ws, der_article):
    titres = ["Date", "Article", "Quantité (unités d'achat)", "Montant payé (F)",
              "Fournisseur", "Mode de paiement", "Remarque",
              "Poste du livre", "Catégorie", "Unité d'achat", "Prix payé par unité (F)",
              "Réf. min (F)", "Réf. max (F)", "Contrôle prix", "ligne"]
    entete(ws, 1, titres, [11, 36, 11, 14, 22, 14, 28, 13, 22, 20, 13, 11, 11, 24, 6])
    A = "Articles!"
    for r in range(2, LIGNES + 2):
        for c in range(1, 8):
            cell = ws.cell(row=r, column=c)
            cell.fill, cell.font, cell.border = FOND_SAISIE, F_SAISIE, CADRE
        ws.cell(row=r, column=1).number_format = FMT_D
        ws.cell(row=r, column=3).number_format = FMT_Q
        ws.cell(row=r, column=4).number_format = FMT_F
        k = f"$O{r}"
        formules = {
            8:  f'=IF($B{r}="","",IF({k}="","article inconnu",INDEX({A}$B:$B,{k})))',
            9:  f'=IF({k}="","",INDEX({A}$C:$C,{k}))',
            10: f'=IF({k}="","",IF(INDEX({A}$D:$D,{k})="","",INDEX({A}$D:$D,{k})))',
            11: f'=IF(OR($C{r}="",$D{r}="",N($C{r})=0),"",$D{r}/$C{r})',
            12: f'=IF({k}="","",IF(INDEX({A}$E:$E,{k})="","",INDEX({A}$E:$E,{k})))',
            13: f'=IF({k}="","",IF(INDEX({A}$F:$F,{k})="","",INDEX({A}$F:$F,{k})))',
            14: (f'=IF(OR($K{r}="",{k}=""),"",IF($L{r}="","pas de prix de référence",'
                 f'IF($K{r}>$M{r}*(1+{A}$K$1),"au-dessus de la fourchette",'
                 f'IF($K{r}<$L{r}*(1-{A}$K$1),"en dessous de la fourchette",'
                 f'"dans la fourchette"))))'),
            15: f'=IF($B{r}="","",IFERROR(MATCH($B{r},{A}$A:$A,0),""))',
        }
        for c, f in formules.items():
            cell = ws.cell(row=r, column=c, value=f)
            cell.fill, cell.font, cell.border = FOND_CALCUL, F_LIEN if c != 11 and c != 14 else F_CALCUL, CADRE
            if c in (11, 12, 13):
                cell.number_format = FMT_F
    ws.column_dimensions["O"].hidden = True
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:N{LIGNES + 1}"

    fin = LIGNES + 1
    dv_art = DataValidation(type="list", formula1=f"=Articles!$A$2:$A${der_article}",
                            allow_blank=True)
    dv_art.error = "Choisir un article de la liste, ou « Autre achat … » avec une remarque."
    dv_art.errorTitle = "Article inconnu"
    dv_date = DataValidation(type="date", operator="between",
                             formula1="DATE(2023,1,1)", formula2="DATE(2035,12,31)")
    dv_date.error, dv_date.errorTitle = "Saisir une date, par exemple 14/10/2026.", "Date"
    dv_qte = DataValidation(type="decimal", operator="greaterThan", formula1="0")
    dv_qte.error, dv_qte.errorTitle = "Une quantité positive, en unités d'achat.", "Quantité"
    dv_mnt = DataValidation(type="whole", operator="greaterThan", formula1="0")
    dv_mnt.error, dv_mnt.errorTitle = "Un montant entier positif, en francs CFA.", "Montant"
    dv_pay = DataValidation(type="list", formula1=f'"{PAIEMENTS}"', allow_blank=True)
    for dv, col in ((dv_date, "A"), (dv_art, "B"), (dv_qte, "C"), (dv_mnt, "D"), (dv_pay, "F")):
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{fin}")
    ws["C1"].comment = Comment(
        "Nombre d'unités d'achat, telles qu'écrites en colonne J (« Unité d'achat ») : "
        "2 bouteilles de 25 kg → 2.", "40_feuille_achats.py")


def plage(col):
    return f"Saisie!${col}$2:${col}${LIGNES + 1}"


def feuille_synthese(ws, articles):
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Articles suivis en priorité — montant acheté, rythme et prix"
    ws["A1"].font = F_TITRE
    ws["A2"], ws["B2"] = "Année", 2026
    ws["A2"].font = F_GRAS
    ws["B2"].font, ws["B2"].fill, ws["B2"].border = F_SAISIE, FOND_SAISIE, CADRE
    ws["C2"] = ("Montants en francs CFA, lus dans « Saisie ». Le rapport aux ventes (question 3) "
                "se fera dans le tableau de bord, qui connaît le chiffre d'affaires.")
    ws["C2"].font = F_NOTE

    titres = (["Article", "Poste"] + [None] * 12 +
              ["Total de l'année", "Nombre d'achats", "Quantité (unités d'achat)",
               "Prix moyen payé par unité", "Premier achat", "Dernier achat",
               "Intervalle moyen (jours)"])
    largeurs = [30, 9] + [10] * 12 + [14, 10, 12, 13, 12, 12, 11]
    entete(ws, 4, [t or "" for t in titres], largeurs)
    for m in range(12):
        c = ws.cell(row=4, column=3 + m, value=f"=DATE($B$2,{m + 1},1)")
        c.number_format = FMT_MOIS

    D, B, A, H, C = plage("D"), plage("B"), plage("A"), plage("H"), plage("C")
    prio = [a for a in articles if a["prio"]]
    for i, a in enumerate(prio):
        r = 5 + i
        ws.cell(row=r, column=1, value=a["art"]).font = F_NORMAL
        ws.cell(row=r, column=2, value=a["poste"]).font = F_NORMAL
        debut, fin = "DATE($B$2,1,1)", "DATE($B$2+1,1,1)"
        for m in range(12):
            col = get_column_letter(3 + m)
            ws.cell(row=r, column=3 + m,
                    value=(f'=SUMIFS({D},{B},$A{r},{A},">="&{col}$4,'
                           f'{A},"<"&EDATE({col}$4,1))'))
        n = f"P{r}"
        ws.cell(row=r, column=15, value=f"=SUM(C{r}:N{r})")
        ws.cell(row=r, column=16, value=f'=COUNTIFS({B},$A{r},{A},">="&{debut},{A},"<"&{fin})')
        ws.cell(row=r, column=17, value=f'=SUMIFS({C},{B},$A{r},{A},">="&{debut},{A},"<"&{fin})')
        ws.cell(row=r, column=18, value=f'=IF(N(Q{r})=0,"",O{r}/Q{r})')
        ws.cell(row=r, column=19, value=(f'=IF({n}=0,"",_xlfn.MINIFS({A},{B},$A{r},'
                                         f'{A},">="&{debut},{A},"<"&{fin}))'))
        ws.cell(row=r, column=20, value=(f'=IF({n}=0,"",_xlfn.MAXIFS({A},{B},$A{r},'
                                         f'{A},">="&{debut},{A},"<"&{fin}))'))
        ws.cell(row=r, column=21, value=f'=IF({n}<2,"",(T{r}-S{r})/({n}-1))')
        for c in range(1, 22):
            cell = ws.cell(row=r, column=c)
            cell.border = CADRE
            if c >= 3:
                cell.font = F_LIEN if c <= 17 or c in (19, 20) else F_CALCUL
            if 3 <= c <= 15 or c == 18:
                cell.number_format = FMT_F
            elif c == 16:
                cell.number_format = FMT_N
            elif c == 17:
                cell.number_format = FMT_Q
            elif c in (19, 20):
                cell.number_format = FMT_D
            elif c == 21:
                cell.number_format = "0.0"
    der = 4 + len(prio)
    t = der + 1
    ws.cell(row=t, column=1, value="Total des articles suivis").font = F_GRAS
    for c in range(3, 16):
        col = get_column_letter(c)
        cell = ws.cell(row=t, column=c, value=f"=SUM({col}5:{col}{der})")
        cell.font, cell.number_format, cell.border = F_GRAS, FMT_F, CADRE
    ws.freeze_panes = "C5"
    ws.cell(row=t + 2, column=1,
            value="Intervalle moyen = jours entre le premier et le dernier achat de l'année, "
                  "divisés par le nombre d'achats moins un.").font = F_NOTE


def feuille_controle(ws):
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Contrôle : total saisi ici, face au livre de dépenses"
    ws["A1"].font = F_TITRE
    ws["A2"], ws["B2"] = "Année", "='Synthèse prioritaires'!B2"
    ws["A2"].font, ws["B2"].font = F_GRAS, F_LIEN
    ws["C2"] = ("Reporter en jaune le total mensuel de chaque colonne du livre de dépenses. "
                "Un écart négatif : des achats du livre n'ont pas été saisis article par article.")
    ws["C2"].font = F_NOTE

    ws.column_dimensions["A"].width = 11
    for i, p in enumerate(POSTES):
        c0 = 2 + 3 * i
        ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c0 + 2)
        h = ws.cell(row=4, column=c0, value=p)
        h.font, h.fill = F_ENTETE, FOND_ENTETE
        h.alignment = Alignment(horizontal="center")
        for j, t in enumerate(("Saisi", "Livre", "Écart")):
            s = ws.cell(row=5, column=c0 + j, value=t)
            s.font, s.fill, s.border = F_ENTETE, FOND_ENTETE, CADRE
            ws.column_dimensions[get_column_letter(c0 + j)].width = 12
    m = ws.cell(row=5, column=1, value="Mois")
    m.font, m.fill, m.border = F_ENTETE, FOND_ENTETE, CADRE

    D, H, A = plage("D"), plage("H"), plage("A")
    for k in range(12):
        r = 6 + k
        mc = ws.cell(row=r, column=1, value=f"=DATE($B$2,{k + 1},1)")
        mc.number_format, mc.font, mc.border = FMT_MOIS, F_CALCUL, CADRE
        for i, p in enumerate(POSTES):
            c0 = 2 + 3 * i
            sl, ll = get_column_letter(c0), get_column_letter(c0 + 1)
            saisi = ws.cell(row=r, column=c0,
                            value=(f'=SUMIFS({D},{H},"{p}",{A},">="&$A{r},'
                                   f'{A},"<"&EDATE($A{r},1))'))
            livre = ws.cell(row=r, column=c0 + 1)
            ecart = ws.cell(row=r, column=c0 + 2,
                            value=f'=IF({ll}{r}="","",{sl}{r}-{ll}{r})')
            saisi.font, ecart.font = F_LIEN, F_CALCUL
            livre.font, livre.fill = F_SAISIE, FOND_SAISIE
            for cell in (saisi, livre, ecart):
                cell.number_format, cell.border = FMT_F, CADRE
    ws.freeze_panes = "B6"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--nomenclature", type=Path,
                    help="defaut : TABOO_NOMENCLATURE dans sources.env")
    ap.add_argument("--sortie", type=Path, default=SORTIE)
    args = ap.parse_args()
    if args.nomenclature is None:
        from sources import source
        args.nomenclature = source("TABOO_NOMENCLATURE")
    if not args.nomenclature.exists():
        raise SystemExit(f"Nomenclature introuvable :\n  {args.nomenclature}")

    articles = lire_nomenclature(args.nomenclature)
    wb = Workbook()
    feuille_lisez_moi(wb.active, articles)
    wb.active.title = "Lisez-moi"
    ws_saisie = wb.create_sheet("Saisie")
    der_article = feuille_articles(wb.create_sheet("Articles"), articles)
    feuille_saisie(ws_saisie, der_article)
    feuille_synthese(wb.create_sheet("Synthèse prioritaires"), articles)
    feuille_controle(wb.create_sheet("Contrôle livre"))
    wb.save(args.sortie)

    prio = sum(a["prio"] for a in articles)
    print(f"  {args.sortie.name}")
    print(f"    {len(articles)} articles au catalogue, dont {prio} suivis en priorité")
    print(f"    {LIGNES} lignes de saisie préparées")
    for p in POSTES:
        print(f"    {p:<13}{sum(a['poste'] == p for a in articles):>4} articles")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
