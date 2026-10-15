#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TABOO — Extraction sans pandas : la meme couche de donnees que
20_extraction_v2.py, en Python standard.

  python 21_extraction_taboo.py                    -> data_v2.json
  python 21_extraction_taboo.py --sortie autre.json

POURQUOI UNE REECRITURE
-----------------------
20_extraction_v2.py repose sur pandas, donc sur numpy et ses
bibliotheques compilees. Sur un poste ou une strategie de controle des
applications bloque ces bibliotheques -- c'est le cas du poste de
controle de gestion -- l'extraction ne tourne tout simplement pas, et
les chiffres de TABOO ne peuvent plus etre mis a jour que d'ailleurs.
Le projet ne doit pas dependre d'un poste ; il ne doit pas non plus
dependre d'une DLL.

Ce fichier ne change AUCUNE regle metier. Il les reprend une a une,
dans le meme ordre, avec les memes invariants. Sur les exports du
26/08/2026, il produit un data_v2.json identique a celui de la version
pandas, hors date de generation : c'est ce qui l'autorise a la
remplacer.

LES REGLES, RAPPELEES
---------------------
    - Journee d'exploitation 14h00 -> 13h59 pour le detail des ventes.
    - Le compte de resultat prend le CA de la Balance par categorie, a sa
      DATE CALENDAIRE : c'est ce que faisait la v2, et le modifier ici
      deplacerait des chiffres deja presentes a la direction.
    - Remise : source de verite = Balance par categorie. Journal CAF
      exclu (rupture de collecte depuis octobre 2024).
    - Achats DRINK = DRINK + MIAMI 228 + PICASSO + GLACONS ; EAT = EAT +
      GAZ ; SMOKE = SMOKE. OPEX 12 postes. CAPEX = EQUIPEMENTS +
      TRAVAUX. CAISSE exclu (fonds de caisse).
    - OPEX et CAPEX ventiles au prorata du CA net du jour.
    - Aucun ratio stocke : des sommes, au grain le plus fin.

CE QUI IMITE PANDAS, ET POURQUOI
--------------------------------
Pour obtenir les memes chiffres, trois comportements de pandas sont
reproduits explicitement : une cellule vide (ou « NA », « NULL »...)
vaut absent ; une ligne dont une cle de regroupement est absente sort
du regroupement ; une somme ignore les absents. Les arrondis sont ceux
de Python, au pair le plus proche, comme ceux de numpy.
"""

import argparse
import csv
import json
import re
import sys
import unicodedata
from collections import OrderedDict
from datetime import datetime, timedelta
from pathlib import Path

from sources import source, trouver

BASE = Path(__file__).resolve().parent

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# Les valeurs que pandas lit comme absentes, par defaut.
ABSENTS = {"", "#N/A", "#N/A N/A", "#NA", "-1.#IND", "-1.#QNAN", "-NaN",
           "-nan", "1.#IND", "1.#QNAN", "<NA>", "N/A", "NA", "NULL", "NaN",
           "None", "n/a", "nan", "null"}

JOURS_SEMAINE = {"lundi", "mardi", "mercredi", "jeudi", "vendredi",
                 "samedi", "dimanche"}

MAPPING_ACHATS = OrderedDict([("DRINK", "DRINK"), ("MIAMI 228", "DRINK"),
                              ("PICASSO", "DRINK"), ("GLACONS", "DRINK"),
                              ("EAT", "EAT"), ("GAZ", "EAT"), ("SMOKE", "SMOKE")])
POSTES_OPEX = ["MONNAIE", "MARKETING", "CACHETS", "CEET/CASH POWER", "TELEPHONIE",
               "INTERNET / TV", "LOYERS", "ADMINISTRATIF", "CONSOMMABLES",
               "ENTRETIEN", "TRANSPORT", "AUTRE"]
POSTES_CAPEX = ["EQUIPEMENTS", "TRAVAUX"]


# ---------------------------------------------------------------------
# Lecture
# ---------------------------------------------------------------------
def val(x):
    """La valeur d'une cellule, ou None si pandas l'aurait lue absente."""
    if x is None:
        return None
    return None if x in ABSENTS else x


def num(s):
    """Nettoyage numerique commun : espaces fines, virgules decimales, #REF!."""
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return float(s)
    s = str(s).replace(" ", "").replace(" ", "").replace(" ", "").replace(",", ".")
    s = s.replace("CFA", "").replace("%", "")
    if "#REF" in s or s in ("", "-"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def z(x):
    """Un absent compte zero dans une somme, comme chez pandas."""
    return 0.0 if x is None else x


def lire_csv(chemin, sep=";", encodage="utf-8-sig"):
    """Les lignes d'un CSV, en dictionnaires, valeurs absentes -> None."""
    with open(chemin, encoding=encodage, newline="") as f:
        lecteur = csv.reader(f, delimiter=sep)
        entete = next(lecteur)
        lignes = []
        for brut in lecteur:
            if not brut:
                continue
            brut = brut + [""] * (len(entete) - len(brut))
            lignes.append({c: val(v) for c, v in zip(entete, brut)})
    return entete, lignes


def date_heure(s, fmt):
    if s is None:
        return None
    try:
        return datetime.strptime(s, fmt)
    except ValueError:
        return None


def jstr(d):
    return d.strftime("%Y-%m-%d")


def cle(s):
    """Normalise un libelle produit pour le rapprochement ventes <-> couts."""
    if s is None:
        return ""
    s = str(s).upper().strip()
    s = "".join(c for c in unicodedata.normalize("NFD", s)
                if unicodedata.category(c) != "Mn")
    return " ".join(s.split())


def regrouper(lignes, cles, mesures, compte=None):
    """groupby(cles).sum(mesures), cles absentes exclues, sortie triee par cle.

    `compte` ajoute le nombre de lignes du groupe (size)."""
    groupes = {}
    for r in lignes:
        k = tuple(r[c] for c in cles)
        if any(x is None for x in k):
            continue
        g = groupes.get(k)
        if g is None:
            g = groupes[k] = {m: 0.0 for m in mesures}
            if compte:
                g[compte] = 0
        for m in mesures:
            g[m] += z(r[m])
        if compte:
            g[compte] += 1
    sortie = []
    for k in sorted(groupes):
        d = dict(zip(cles, k))
        d.update(groupes[k])
        sortie.append(d)
    return sortie


MOIS_FR = ["JANVIER", "FEVRIER", "MARS", "AVRIL", "MAI", "JUIN", "JUILLET",
           "AOUT", "SEPTEMBRE", "OCTOBRE", "NOVEMBRE", "DECEMBRE"]


def sans_accents(s):
    return "".join(c for c in unicodedata.normalize("NFD", str(s))
                   if unicodedata.category(c) != "Mn")


# ---------------------------------------------------------------------
# Rapports mensuels du logiciel de caisse
# ---------------------------------------------------------------------
def lots_mensuels():
    """Les dossiers de rapports mensuels configures, chacun lu en entier.

    Un dossier qui contient lui-meme un « journal de vente » est un lot ;
    sinon, chacun de ses sous-dossiers qui en contient un en est un."""
    from sources import configure
    if not configure("TABOO_RAPPORTS_MENSUELS"):
        return []
    racine = source("TABOO_RAPPORTS_MENSUELS", dossier=True)
    a_un_journal = lambda d: any("journal de vente" in sans_accents(p.name).lower()
                                 for p in d.glob("*.csv"))
    dossiers = [racine] if a_un_journal(racine) else sorted(
        d for d in racine.iterdir() if d.is_dir() and a_un_journal(d))
    lots = []
    for d in dossiers:
        lot = {"dossier": d}
        for nom, fragment in (("ventes", "journal de vente"), ("caf", "journal CAF"),
                              ("depenses", "états des dépenses"),
                              ("reglements", "états des règlements"),
                              ("famille", "états par famille")):
            lot[nom] = lire_csv(trouver(d, fragment))[1]
        lots.append(lot)
    return lots


def lire_livre_classeur(chemin):
    """Le livre de depenses de l'onglet Detail_Depenses du classeur.

    Memes colonnes que le CSV historique ; les lignes sont rendues sous la
    meme forme que lire_csv(), pour que la suite ne voie pas la difference."""
    import openpyxl
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    # Le classeur a DEUX onglets « detail depenses ». DETAIL_DEPENSES est une
    # table derivee, ancienne, sans colonne JOURS ; le livre, Détail_Dépenses,
    # a JOURS, DATES et les postes. On le reconnait a son contenu, pas a son
    # orthographe, qui ne differe que par les accents.
    onglet, lignes = None, None
    for n in wb.sheetnames:
        if sans_accents(n).strip().upper() != "DETAIL_DEPENSES":
            continue
        candidat = list(wb[n].iter_rows(values_only=True))
        tete = {str(c).strip() for c in candidat[0] if c} if candidat else set()
        if {"JOURS", "DATES", "EAT", "DRINK"} <= tete:
            onglet, lignes = n, candidat
            break
    if onglet is None:
        raise SystemExit(f"Aucun onglet Detail_Depenses avec JOURS, DATES et les postes dans {chemin}")
    entete = [str(c).strip() if c is not None else "" for c in lignes[0]]
    sortie = []
    for brut in lignes[1:]:
        r = {}
        for c, x in zip(entete, brut):
            if not c:
                continue
            if hasattr(x, "strftime"):
                x = x.strftime("%d/%m/%Y")
            elif isinstance(x, float) and x.is_integer():
                x = str(int(x))
            r[c] = val(str(x)) if x is not None else None
        sortie.append(r)
    return onglet, sortie


def onglet_du_mois(chemin, annee, mois):
    """Les lignes jour par jour de l'onglet « SEPTEMBRE 26 » du classeur."""
    import openpyxl
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    voulu = f"{MOIS_FR[mois - 1]} {annee % 100:02d}"
    nom = next((n for n in wb.sheetnames if sans_accents(n).strip().upper() == voulu), None)
    if nom is None:
        return None, {}
    lignes = list(wb[nom].iter_rows(values_only=True))
    entete = [sans_accents(c).strip().upper() if c else "" for c in lignes[0]]
    jours = {}
    for r in lignes[1:]:
        if len(r) > 1 and hasattr(r[1], "strftime"):
            jours[r[1].date()] = {c: (x if isinstance(x, (int, float)) else 0)
                                  for c, x in zip(entete, r) if c}
    return nom, jours


_alertes = []


def invariant(ok, message):
    if not ok:
        raise SystemExit(f"\nINCOHERENCE : {message}\n  Arret : mieux vaut "
                         f"pas de dashboard qu'un dashboard faux.\n")


# ---------------------------------------------------------------------
# Extraction
# ---------------------------------------------------------------------
def extraire(comme_v2=False):
    BQ = source("TABOO_EXPORTS_POS", dossier=True)
    F_VENTES = trouver(BQ, "Journal_Vente", colonnes=("Date", "Heure"))
    F_FAMILLE = trouver(BQ, "Balance par catégorie")
    F_OFFERTS = trouver(BQ, "Articles offerts")
    F_PANIER = trouver(BQ, "Panier moyen horaire")
    F_REGL = trouver(BQ, "Balance par règlement")
    F_UTIL = trouver(BQ, "Balance par utilisateur")
    F_CAGLOBAL = trouver(BQ, "Chiffre d'affaire globale")
    F_INVENT = trouver(BQ, "INVENTAIRE")
    F_DEPENSES = source("TABOO_DEPENSES")
    F_COUTS = source("COUTS_DE_REVIENT")

    print("=" * 70)
    print("TABOO — extraction (Python standard)")
    print("=" * 70)

    # =================================================================
    # 1. VENTES DETAILLEES — grain article, journee d'exploitation 14h
    # =================================================================
    print("\n[1/9] Journal de vente detaille...")
    # Prix de vente constates par article : les articles offerts du rapport
    # mensuel arrivent avec un prix nul, on les valorise au prix habituel.
    prix_vus = {}

    def ligne_vente(r):
        if r["Type"] == "TOTAL":
            return None
        ds = date_heure(r["Date"], "%d/%m/%Y %H:%M:%S")
        hs = date_heure((r["Heure"] or "nan")[:8], "%H:%M:%S")
        if ds is None or hs is None:
            return None
        h = hs.hour
        jour = ds.date() if h >= 14 else (ds - timedelta(days=1)).date()
        caf, remise, pu = num(r["CAF"]), num(r["Remise"]), num(r["Pu"])
        if pu and r["Articles"] is not None:
            p = prix_vus.setdefault(r["Articles"], {})
            p[pu] = p.get(pu, 0) + 1
        return {"Type": r["Type"], "Categorie": r["Categorie"],
                "Articles": r["Articles"], "Qté": num(r["Qté"]),
                "Remise": remise, "CAF": caf,
                # CAF est le CA NET (apres remise). Le brut se reconstitue.
                "ca_brut_ligne": None if caf is None else caf + z(remise),
                "jour": jour, "h": h, "ordre_nuit": h - 14 if h >= 14 else h + 10,
                "mois": jour.strftime("%Y-%m"), "calendaire": ds.date()}

    _, brut = lire_csv(F_VENTES)
    v = [x for x in map(ligne_vente, brut) if x is not None]
    calendaires_hist = {r["calendaire"] for r in v}

    ca_net_ventes = sum(z(r["CAF"]) for r in v)
    invariant(len(v) > 0, "journal de vente vide")
    invariant(ca_net_ventes > 0, f"CA net nul : {ca_net_ventes:,.0f}")
    jours_v = sorted({r["jour"] for r in v})
    print(f"      {len(v):,} lignes · CA net {ca_net_ventes:,.0f} F")
    print(f"      periode {jours_v[0]} -> {jours_v[-1]}")

    # =================================================================
    # 2. BALANCE PAR CATEGORIE — source de verite brut / remise / net
    # =================================================================
    print("\n[2/9] Balance par categorie (source de verite remise)...")
    _, brut = lire_csv(F_FAMILLE)
    e = []
    for r in brut:
        dc = date_heure(r["Date"], "%d/%m/%Y %H:%M:%S")
        if dc is None or r["Type"] not in ("DRINK", "EAT", "SMOKE"):
            continue
        e.append({"jour": dc, "Type": r["Type"], "Qté": num(r["Qté"]),
                  "offert": num(r["offert"]), "TTC": num(r["TTC"]),
                  "Total Remise": num(r["Total Remise"]), "Total HT": num(r["Total HT"])})
    brut_t = sum(z(r["TTC"]) for r in e)
    remise_t = sum(z(r["Total Remise"]) for r in e)
    net = sum(z(r["Total HT"]) for r in e)
    invariant(abs((brut_t - remise_t) - net) < 1,
              f"brut - remise != net ({brut_t:,.0f} - {remise_t:,.0f} != {net:,.0f})")
    invariant(abs(net - ca_net_ventes) < 1,
              f"les deux exports POS divergent sur le CA net : "
              f"journal {ca_net_ventes:,.0f} vs balance {net:,.0f}")
    print(f"      brut {brut_t:,.0f} · remise {remise_t:,.0f} · net {net:,.0f} F")

    # =================================================================
    # 2 bis. RAPPORTS MENSUELS — les mois venus apres l'export historique
    # =================================================================
    # Le rapport mensuel n'a pas de Balance par categorie DATEE (son « etats
    # par famille » est un cumul du mois). Le CA du compte de resultat s'y
    # prend donc dans le journal de vente, a la date calendaire : verifie sur
    # tout l'historique, le CA net du journal ainsi regroupe egale celui de
    # la Balance sur 3 579 jours x activite sur 3 579. Les quantites vendues
    # et offertes se separent par le CA de la ligne (nul = offert).
    hist_fin = jours_v[-1]
    lots = [] if comme_v2 else lots_mensuels()
    couverts = set(calendaires_hist)
    infos_lots = []
    for lot in lots:
        lignes = [x for x in map(ligne_vente, lot["ventes"]) if x is not None]
        cal = {r["calendaire"] for r in lignes}
        deja = sorted(cal & couverts)
        invariant(not deja,
                  f"{lot['dossier'].name} recouvre des jours deja charges "
                  f"({deja[0] if deja else ''}...) : un jour compte deux fois fausserait tout")
        couverts |= cal
        lot["lignes"], lot["calendaires"] = lignes, cal
    for lot in lots:
        # Une ligne vendue le 01/09 avant 14 h appartient a la journee du
        # 31/08. Si aucun lot ne couvre ce jour-la, la garder ferait une
        # journee faite de sa seule fin de nuit : on l'ecarte, et on le dit.
        gardees = [r for r in lot["lignes"] if r["jour"] in couverts]
        ecartees = len(lot["lignes"]) - len(gardees)
        v.extend(gardees)
        for r in lot["lignes"]:
            if r["Type"] not in ("DRINK", "EAT", "SMOKE"):
                continue
            offert = z(r["CAF"]) == 0
            e.append({"jour": datetime.combine(r["calendaire"], datetime.min.time()),
                      "Type": r["Type"],
                      "Qté": 0.0 if offert else r["Qté"], "offert": r["Qté"] if offert else 0.0,
                      "TTC": r["ca_brut_ligne"], "Total Remise": r["Remise"],
                      "Total HT": r["CAF"]})
        debut, fin = min(lot["calendaires"]), max(lot["calendaires"])
        ca_lot = sum(z(r["CAF"]) for r in lot["lignes"])
        infos_lots.append({"lot": lot, "debut": debut, "fin": fin, "ca": ca_lot,
                           "ecartees": ecartees})
        print(f"      + {lot['dossier'].name} : {debut} -> {fin} · "
              f"{len(lot['lignes']):,} lignes · CA net {ca_lot:,.0f} F"
              + (f" · {ecartees} lignes d'une journee non couverte ecartees" if ecartees else ""))
    jours_v = sorted({r["jour"] for r in v})

    # =================================================================
    # 3. DEPENSES
    # =================================================================
    print("\n[3/9] Depenses...")
    from sources import configure
    if configure("TABOO_ANALYSE_MENSUELLE") and not comme_v2:
        # Le classeur du controle de gestion fait foi : son livre contient
        # des saisies posterieures a l'export CSV (104 cellules differaient
        # jusqu'au 23/08/2026 lors de la bascule).
        F_CLASSEUR = source("TABOO_ANALYSE_MENSUELLE")
        onglet_livre, brut = lire_livre_classeur(F_CLASSEUR)
        print(f"      livre : {F_CLASSEUR.name} / {onglet_livre.strip()}")
    else:
        F_CLASSEUR = None
        _, brut = lire_csv(F_DEPENSES, sep=",", encodage="utf-8")
    # Borne haute deduite des ventes : le livre est tenu a la main et contient
    # des lignes previsionnelles au-dela du dernier jour vendu. Au-dela de
    # l'export historique, on ne garde que les jours couverts par un rapport
    # mensuel : un achat sans les ventes du meme jour fausserait la marge.
    jours_lots = set().union(*(l["calendaires"] for l in lots)) if lots else set()
    dep = []
    for r in brut:
        if r.get("DATES") is None:
            continue
        if (r.get("JOURS") or "").lower() not in JOURS_SEMAINE:
            continue
        j = date_heure(r["DATES"], "%d/%m/%Y")
        if j is None or not (j.date() <= hist_fin or j.date() in jours_lots):
            continue
        j = j.date()
        for poste, act in MAPPING_ACHATS.items():
            m = num(r.get(poste))
            if m is not None and m != 0:
                dep.append({"jour": j, "nature": "ACHATS EXTERNES", "activite": act,
                            "poste": poste, "montant": m})
        for nature, postes in (("OPEX", POSTES_OPEX), ("CAPEX", POSTES_CAPEX)):
            for poste in postes:
                m = num(r.get(poste))
                if m is not None and m != 0:
                    dep.append({"jour": j, "nature": nature, "activite": None,
                                "poste": poste, "montant": m})
    invariant(len(dep) > 0, "aucune depense retenue : colonnes renommees ?")
    print(f"      {len(dep):,} lignes · total {sum(r['montant'] for r in dep):,.0f} F")

    # =================================================================
    # 4. RESULTAT JOUR x ACTIVITE — union ventes + achats
    # =================================================================
    print("\n[4/9] Compte de resultat au grain jour x activite...")
    ca_ja = {}
    for g in regrouper(e, ["jour", "Type"], ["TTC", "Total Remise", "Total HT", "Qté", "offert"]):
        ca_ja[(g["jour"].date(), g["Type"])] = g
    ach_ja = {}
    for g in regrouper([r for r in dep if r["nature"] == "ACHATS EXTERNES"],
                       ["jour", "activite"], ["montant"]):
        ach_ja[(g["jour"], g["activite"])] = g["montant"]

    # FULL OUTER : un achat un jour sans vente ne doit pas disparaitre.
    res = []
    for k in sorted(set(ca_ja) | set(ach_ja)):
        g = ca_ja.get(k, {})
        res.append({"jour": k[0], "activite": k[1],
                    "ca_brut": g.get("TTC", 0.0), "remise": g.get("Total Remise", 0.0),
                    "ca_net": g.get("Total HT", 0.0), "qv": g.get("Qté", 0.0),
                    "qo": g.get("offert", 0.0), "achats": ach_ja.get(k, 0.0)})

    # OPEX / CAPEX ventiles au prorata du CA net du jour.
    def par_jour(nature):
        t = {}
        for r in dep:
            if r["nature"] == nature:
                t[r["jour"]] = t.get(r["jour"], 0.0) + r["montant"]
        return t
    opex_j, capex_j = par_jour("OPEX"), par_jour("CAPEX")
    canet_j = {}
    for r in res:
        canet_j[r["jour"]] = canet_j.get(r["jour"], 0.0) + r["ca_net"]
    for r in res:
        c = canet_j[r["jour"]]
        part = r["ca_net"] / c if c > 0 else 0
        r["opex"] = float(round(opex_j.get(r["jour"], 0.0) * part))
        r["capex"] = float(round(capex_j.get(r["jour"], 0.0) * part))
        r["mb"] = r["ca_net"] - r["achats"]
        r["re"] = r["mb"] - r["opex"]
        r["rn"] = r["re"] - r["capex"]

    achats_src = sum(r["montant"] for r in dep if r["nature"] == "ACHATS EXTERNES")
    invariant(abs(sum(r["achats"] for r in res) - achats_src) < 1,
              "achats perdus dans la jointure (regression du FULL OUTER JOIN)")
    tot_cn = sum(r["ca_net"] for r in res)
    print(f"      marge brute {sum(r['mb'] for r in res):,.0f} F "
          f"({100*sum(r['mb'] for r in res)/tot_cn:.1f} %)")
    print(f"      resultat net {sum(r['rn'] for r in res):,.0f} F "
          f"({100*sum(r['rn'] for r in res)/tot_cn:.1f} %)")

    # =================================================================
    # 5. ARTICLES — catalogue complet
    # =================================================================
    print("\n[5/9] Catalogue articles complet...")
    art = regrouper(v, ["Type", "Categorie", "Articles"],
                    ["Qté", "CAF", "ca_brut_ligne", "Remise"], compte="nl")
    art.sort(key=lambda a: -a["CAF"])
    for a in art:
        a["cle"] = cle(a["Articles"])
    print(f"      {len(art):,} articles distincts")
    total_art = sum(a["CAF"] for a in art)
    cumul, n80 = 0.0, 0
    for a in art:
        cumul += a["CAF"]
        if cumul / total_art <= 0.8:
            n80 += 1
    n80 += 1
    print(f"      Pareto : {n80} articles ({100*n80/len(art):.1f} %) font 80 % du CA")

    # =================================================================
    # 6. COUTS DE REVIENT
    # =================================================================
    print("\n[6/9] Couts de revient...")
    _, brut = lire_csv(F_COUTS, encodage="utf-8")
    cr = {}
    for r in brut:
        cout = num(r["cout_revient"])
        if cout is None:
            continue
        k = cle(r["produit"])
        if k in cr:
            continue
        cr[k] = {"type_source": r["type_source"], "cout_revient": cout,
                 "prix_vente": num(r["prix_vente"]), "statut_cout": r["statut_cout"]}
    apparies = [a for a in art if a["cle"] in cr]
    ca_couvert = sum(a["CAF"] for a in apparies)
    couv = 100 * ca_couvert / total_art
    print(f"      {len(cr):,} produits references · {len(apparies):,} apparies "
          f"aux ventes ({couv:.1f} % du CA couvert)")
    if couv < 20:
        _alertes.append(f"Couts de revient : seulement {couv:.1f} % du CA est couvert. "
                        f"La marge theorique ne porte que sur ce perimetre.")
    couts_articles = []
    for a in apparies:
        if not a["Qté"] > 0:
            continue
        c = cr[a["cle"]]
        cout_total = c["cout_revient"] * a["Qté"]
        couts_articles.append({
            "t": a["Type"], "c": a["Categorie"], "a": a["Articles"], "src": c["type_source"],
            "q": int(a["Qté"]), "cn": round(a["CAF"]), "cb": round(a["ca_brut_ligne"]),
            "rm": round(a["Remise"]), "cr": round(c["cout_revient"], 2),
            "pv": round(c["prix_vente"]) if c["prix_vente"] is not None else None,
            "ct": round(cout_total), "mt": round(a["CAF"] - cout_total),
            "st": c["statut_cout"]})

    # =================================================================
    # 7. INVENTAIRE
    # =================================================================
    print("\n[7/9] Inventaire...")
    entete, brut = lire_csv(F_INVENT)
    inv = []
    for r in brut:
        r = {k.strip(): x for k, x in r.items()}
        if r["Articles"] is None:
            continue
        for c in ("Qté théorique", "Stock Alerte", "Qté physique", "Ecart"):
            r[c] = num(r[c])
        inv.append(r)
    compte = sum(1 for r in inv if r["Qté physique"] is not None and r["Qté physique"] > 0)
    inventaire_compte = compte > 0
    if not inventaire_compte:
        _alertes.append(
            "Inventaire : la quantite physique est a 0 sur les "
            f"{len(inv)} articles — le comptage n'a jamais ete saisi dans le POS. "
            "L'ecart affiche n'est donc PAS une demarque : il vaut mecaniquement "
            "l'oppose du stock theorique. Seuls le stock theorique et le seuil "
            "d'alerte sont exploitables.")
    print(f"      {len(inv):,} articles · comptage physique saisi sur {compte}")

    def i0(x):
        return 0 if x is None else int(x)

    inventaire, n_alerte, val_stock = [], 0, 0.0
    for r in inv:
        cout = cr.get(cle(r["Articles"]), {}).get("cout_revient")
        qt, sa = r["Qté théorique"], r["Stock Alerte"]
        vs = qt * cout if (qt is not None and cout is not None) else None
        alerte = (sa is not None and sa > 0) and (qt is not None and qt <= sa)
        n_alerte += alerte
        val_stock += z(vs)
        inventaire.append({"cat": r["Catégorie"], "a": r["Articles"], "qt": i0(qt),
                           "sa": i0(sa), "qp": i0(r["Qté physique"]), "ec": i0(r["Ecart"]),
                           "vs": round(vs) if vs is not None else None, "al": bool(alerte)})
    print(f"      {n_alerte} articles sous seuil d'alerte · stock valorise "
          f"{val_stock:,.0f} F (sur le perimetre a cout connu)")

    # =================================================================
    # 8. AUTRES SOURCES POS
    # =================================================================
    print("\n[8/9] Reglements, horaire, caissiers, offerts, totaux POS...")

    # --- Reglements (cumul, aucune date dans la source)
    _, brut = lire_csv(F_REGL)
    noms = list(brut[0].keys()) if brut else []
    rg = [{"libelle": r[noms[0]], "nombre": num(r[noms[1]]), "montant": num(r[noms[2]])}
          for r in brut]
    MOYENS = ["CB", "CHEQUES", "ESPECES", "FLOOZ", "GOZEM", "TMONEY"]
    regl = [r for r in rg if r["libelle"] in MOYENS]
    cred = [r for r in rg if r["libelle"] is not None
            and re.search("crédit|credit", r["libelle"], re.I)]
    regl.append({"libelle": "Crédit",
                 "nombre": float(cred[0]["nombre"]) if cred else 0.0,
                 "montant": float(cred[0]["montant"]) if cred else 0.0})

    # --- Profil horaire de reference (source dediee, sans date)
    _, brut = lire_csv(F_PANIER)
    noms = list(brut[0].keys()) if brut else []
    ph = []
    for r in brut:
        x = [r[n] for n in noms]
        p = {"heure": x[0], "ca": num(x[1]), "ventes": num(x[2]),
             "vendeurs": num(x[3]), "panier_moyen": num(x[4])}
        if p["ca"] is None:
            continue
        m = re.match(r"^(\d+)", p["heure"] or "")
        hd = float(m.group(1)) if m else None
        p["heure_debut"] = hd
        p["ordre_nuit"] = (hd - 14 if hd >= 14 else hd + 10) if hd is not None else None
        ph.append(p)
    ph.sort(key=lambda p: p["ordre_nuit"])
    panier_horaire = [{k: p[k] for k in ("heure", "heure_debut", "ordre_nuit", "ca",
                                         "ventes", "vendeurs", "panier_moyen")} for p in ph]

    # --- Caissiers (export en blocs cle/valeur, sans date)
    _, brut = lire_csv(F_UTIL)
    noms = list(brut[0].keys()) if brut else []
    blocs, cur = [], None
    for r in brut:
        lib = str(r[noms[0]]).strip() if r[noms[0]] is not None else "nan"
        valeur = r[noms[1]]
        if lib == "Utilisateur":
            if cur:
                blocs.append(cur)
            cur = {"utilisateur": str(valeur).strip() if valeur is not None else "nan"}
        elif cur is not None and lib and lib != "nan":
            cur[lib] = valeur
    if cur:
        blocs.append(cur)
    colonnes = []
    for b in blocs:
        for k in b:
            if k not in colonnes:
                colonnes.append(k)
    numeriques = {"Nbre produit", "Nbre de ticket", "Nbre panier", "Panier Moyen",
                  "Total HT", "Total TTC", "Total règlement"}
    caissiers = [{k: (num(b.get(k)) if k in numeriques else b.get(k)) for k in colonnes}
                 for b in blocs]

    # --- Articles offerts
    _, brut = lire_csv(F_OFFERTS)
    off = []
    for r in brut:
        d = date_heure(r["Offert le"], "%d/%m/%Y %H:%M:%S")
        if d is None:
            continue
        off.append({"jour": d, "Caissier": r["Caissier"], "Catégorie": r["Catégorie"],
                    "Article": r["Article"], "Quantité": num(r["Quantité"]),
                    "Val vente": num(r["Val vente"])})

    # Les rapports mensuels n'ont pas d'export « Articles offerts » : les
    # offerts y sont les lignes du journal a CA nul, avec un prix nul. On les
    # valorise au prix de vente le plus frequent de l'article, et le caissier
    # n'est pas connu -- ce qui se dit plutot que de s'inventer.
    def prix_habituel(article):
        p = prix_vus.get(article)
        return max(p, key=lambda x: (p[x], x)) if p else None

    for info in infos_lots:
        sans_prix, valeur = 0, 0.0
        for r in info["lot"]["lignes"]:
            if z(r["CAF"]) != 0:
                continue
            pu = prix_habituel(r["Articles"])
            sans_prix += pu is None
            montant = z(pu) * z(r["Qté"])
            valeur += montant
            off.append({"jour": datetime.combine(r["calendaire"], datetime.min.time()),
                        "Caissier": "(non détaillé)", "Catégorie": r["Categorie"],
                        "Article": r["Articles"], "Quantité": r["Qté"], "Val vente": montant})
        info["offert_valeur"], info["offert_sans_prix"] = valeur, sans_prix

    # --- Totaux POS de controle (nombre de paniers, panier moyen)
    with open(F_CAGLOBAL, encoding="utf-8-sig", newline="") as f:
        cg = []
        for x in csv.reader(f, delimiter=";"):
            x = x + [""] * (2 - len(x))
            lib = val(x[0])
            lib = "nan" if lib is None else str(lib).strip().rstrip(":").strip()
            cg.append((lib, val(x[1])))

    def cg_val(motif):
        for lib, valeur in cg:
            if re.search(motif, lib, re.I):
                return num(valeur)
        return None

    nb_panier = cg_val(r"^Nbre panier")
    panier_moyen = cg_val(r"^Panier Moyen")
    nombre_vendu = cg_val(r"^Nombre vendu")
    remise_pos = cg_val(r"^Remise")
    total_offert = cg_val(r"^Total Offert")
    ca_ht_hors_offert = cg_val(r"CA H\.T Hors offert")
    ca_offert = cg_val(r"dont CA HT offert")
    print(f"      {nb_panier:,.0f} paniers · panier moyen {panier_moyen:,.0f} F")

    # =================================================================
    # 9. RECONCILIATION — les ecarts documentes, affiches et non caches
    # =================================================================
    print("\n[9/9] Reconciliation...")
    total_regl = sum(z(r["montant"]) for r in regl)
    credit_montant = regl[-1]["montant"]
    reconciliation = [
        {"controle": "Coherence interne de la source de verite",
         "attendu": round(brut_t - remise_t), "obtenu": round(net),
         "regle": "CA brut − remise = CA net",
         "statut": "OK" if abs((brut_t - remise_t) - net) < 1 else "ALERTE",
         "explication": "Verifie que la Balance par categorie est coherente avec "
                        "elle-meme. Tout ecart invaliderait la source de verite de la remise."},
        {"controle": "CA net : journal de vente vs balance par categorie",
         "attendu": round(net), "obtenu": round(ca_net_ventes),
         "regle": "Deux rapports POS independants, meme CA net",
         "statut": "OK" if abs(net - ca_net_ventes) < 1 else "ALERTE",
         "explication": "Les deux exports concordent a l'unite pres. C'est le "
                        "controle le plus fort dont on dispose sur l'ingestion."},
        {"controle": "CA global POS vs balance par categorie",
         "attendu": round(ca_ht_hors_offert), "obtenu": round(net),
         "regle": "Ecart structurel connu",
         "statut": "ECART CONNU",
         # La v2 remplacait les virgules de TOUTE la phrase, pas seulement
         # celles des milliers : reproduit tel quel, pour un fichier identique.
         "explication": "Ecart de {:,.0f} F. Les deux rapports POS ne couvrent pas "
                        "le meme perimetre d'articles. Documente dans "
                        "06_reconciliation.sql ; surveiller sa variation, pas son "
                        "niveau.".format(ca_ht_hors_offert - net).replace(",", " ")},
        {"controle": "CA net vs total des reglements encaisses",
         "attendu": round(net), "obtenu": round(total_regl),
         "regle": "Ecart structurel connu",
         "statut": "ECART CONNU",
         "explication": "Ecart de {:,.0f} F, dont {:,.0f} F de ventes a credit. "
                        "Le solde correspond aux ecarts d'encaissement.".format(
                            net - total_regl, credit_montant).replace(",", " ")},
        {"controle": "Remise : Journal CAF vs balance par categorie",
         "attendu": round(remise_t), "obtenu": None,
         "regle": "Journal CAF ECARTE de la chaine",
         "statut": "SOURCE ECARTEE",
         "explication": "Le Journal CAF sous-evalue la remise de 90 a 99 % depuis "
                        "octobre 2024 (rupture de collecte cote POS). C'est une "
                        "anomalie de l'editeur, pas un ecart metier : la chaine "
                        "n'utilise que la Balance par categorie."},
        {"controle": "Inventaire : comptage physique saisi",
         "attendu": len(inv), "obtenu": compte,
         "regle": "Chaque article doit avoir un comptage",
         "statut": "OK" if inventaire_compte else "DONNEE MANQUANTE",
         "explication": "Aucun comptage physique n'a ete saisi. L'ecart "
                        "d'inventaire n'est donc pas exploitable ; le stock "
                        "theorique et les seuils d'alerte le sont."},
    ]

    # --- Rapports mensuels : la caisse confrontee a l'analyse mensuelle.
    # Le classeur du controle de gestion fait foi ; la caisse, elle, peut
    # avoir ete corrigee apres coup. Chaque ecart est donc chiffre et
    # decompose, jamais absorbe en silence.
    POS_VERS_CLASSEUR = {"ESPECES": "CASH", "CB": "TPE", "TMONEY": "MIXX",
                         "FLOOZ": "FLOOZ", "GOZEM": "GOZEM"}
    f0 = lambda x: "{:,.0f}".format(x).replace(",", " ")
    for info in infos_lots:
        lot, fin = info["lot"], info["fin"]
        mois = f"{MOIS_FR[fin.month - 1].lower()} {fin.year}"
        fam_net = sum(z(num(r["Total HT"])) for r in lot["famille"])
        reconciliation.append({
            "controle": f"CA net {mois} : journal de vente vs etats par famille",
            "attendu": round(fam_net), "obtenu": round(info["ca"]),
            "regle": "Deux rapports du meme mois, meme CA net",
            "statut": "OK" if abs(fam_net - info["ca"]) < 1 else "ALERTE",
            "explication": "Le rapport mensuel n'a pas de balance datee : le CA du "
                           "compte de resultat y vient du journal de vente. Ce controle "
                           "verifie que le journal porte bien tout le CA du mois."})
        caf_ht = sum(z(num(r["HT"])) for r in lot["caf"])
        rem_ticket = sum(z(num(r["Remise"])) for r in lot["caf"])
        ca_offert = sum(z(num(r["CA offert"])) for r in lot["caf"])
        if F_CLASSEUR is None:
            continue
        onglet, ref = onglet_du_mois(F_CLASSEUR, fin.year, fin.month)
        ref = {d: x for d, x in ref.items() if d in lot["calendaires"]}
        if not ref:
            _alertes.append(f"Rapprochement {mois} impossible : aucun onglet "
                            f"« {MOIS_FR[fin.month - 1]} {fin.year % 100:02d} » dans le classeur.")
            continue
        # CA : l'ecart se decompose en remises sur ticket (journal CAF) et en
        # difference entre journal de vente et journal CAF.
        ca_ref = sum(x.get("CA", 0) for x in ref.values())
        ecart = info["ca"] - ca_ref
        hors_caf = info["ca"] - caf_ht
        residu = ecart - rem_ticket - hors_caf
        reconciliation.append({
            "controle": f"CA {mois} : caisse vs analyse mensuelle",
            "attendu": round(ca_ref), "obtenu": round(info["ca"]),
            "regle": "L'analyse mensuelle fait foi",
            "statut": "ECART EXPLIQUE" if abs(residu) < 1 else "ECART A EXPLIQUER",
            "explication": f"La caisse porte {f0(ecart)} F de plus que l'analyse ({onglet.strip()}). "
                           f"Dont {f0(rem_ticket)} F de remises accordees sur le ticket entier, "
                           f"que le detail par article ne porte pas (journal CAF), et "
                           f"{f0(hors_caf)} F d'ecart entre le journal de vente et le journal CAF."
                           + (f" Reste {f0(residu)} F sans explication." if abs(residu) >= 1 else
                              " L'ecart est entierement explique.")})
        # Encaissements, moyen par moyen.
        pos = {}
        for r in lot["reglements"]:
            pos[r["Moyen de paiement"]] = pos.get(r["Moyen de paiement"], 0) + z(num(r["Montant payé"]))
        lignes_moyens, pos_t, ref_t, alerte = [], 0.0, 0.0, False
        for moyen, col in POS_VERS_CLASSEUR.items():
            a, b = pos.get(moyen, 0.0), sum(x.get(col, 0) for x in ref.values())
            pos_t += a
            ref_t += b
            if abs(a - b) >= 1:
                alerte = True
                lignes_moyens.append(f"{moyen.lower()} : caisse {f0(a)} F, analyse {f0(b)} F "
                                     f"(ecart {f0(a - b)} F)")
        autres = {k: x for k, x in pos.items() if k not in POS_VERS_CLASSEUR}
        reconciliation.append({
            "controle": f"Encaissements {mois} : caisse vs analyse mensuelle",
            "attendu": round(ref_t), "obtenu": round(pos_t),
            "regle": "Meme montant par moyen de paiement",
            "statut": "ALERTE" if alerte else "OK",
            "explication": ("Ecarts : " + " ; ".join(lignes_moyens) + ". "
                            if lignes_moyens else "Tous les moyens concordent. ")
                           + "Les especes de la caisse se comparent a la colonne CASH, la "
                             "carte a TPE, TMONEY a MIXX."
                           + (" Moyens de la caisse sans equivalent : "
                              + ", ".join(f"{k} {f0(x)} F" for k, x in autres.items()) + "."
                              if autres else "")})
        # Depenses : la caisse ne voit que les sorties du tiroir.
        sorties = sum(z(num(r["Montant"])) for r in lot["depenses"])
        livre = sum(r["montant"] for r in dep if r["jour"] in lot["calendaires"])
        par_jour_caisse, par_jour_livre = {}, {}
        for r in lot["depenses"]:
            d = date_heure(r["Date"], "%d/%m/%Y %H:%M:%S")
            if d:
                par_jour_caisse[d.date()] = par_jour_caisse.get(d.date(), 0) + z(num(r["Montant"]))
        for r in dep:
            if r["jour"] in lot["calendaires"]:
                par_jour_livre[r["jour"]] = par_jour_livre.get(r["jour"], 0) + r["montant"]
        jours_sup = sorted(d for d, x in par_jour_caisse.items() if x > par_jour_livre.get(d, 0) + 0.5)
        reconciliation.append({
            "controle": f"Depenses {mois} : sorties de caisse vs livre de depenses",
            "attendu": round(livre), "obtenu": round(sorties),
            "regle": "La caisse ne voit que les sorties du tiroir",
            "statut": "ALERTE" if jours_sup else "ECART CONNU",
            "explication": f"La caisse enregistre {f0(sorties)} F de sorties, soit "
                           f"{100 * sorties / livre:.0f} % du livre ({f0(livre)} F) : les "
                           "reglements fournisseurs par virement ou mobile money n'y passent pas."
                           + (f" Mais {len(jours_sup)} jour(s) ou la caisse sort PLUS que le livre "
                              f"n'en note : {', '.join(d.strftime('%d/%m') for d in jours_sup[:8])}"
                              + ("..." if len(jours_sup) > 8 else "") + "."
                              if jours_sup else "")})
        _alertes.append(
            f"{mois.capitalize()} : les articles offerts viennent du journal de vente (lignes a "
            f"CA nul), valorises au prix de vente habituel de l'article — "
            f"{f0(info['offert_valeur'])} F, contre {f0(ca_offert)} F de « CA offert » dans le "
            "journal CAF. Le caissier qui offre n'y est pas connu."
            + (f" {info['offert_sans_prix']} ligne(s) sans prix connu valent 0."
               if info["offert_sans_prix"] else ""))

    # --- Ce qui manque, et ce qui n'est pas mis a jour.
    if infos_lots:
        premier = min(i["debut"] for i in infos_lots)
        if (premier - hist_fin).days > 1:
            _alertes.append(
                f"Aucune donnee du {(hist_fin + timedelta(days=1)):%d/%m/%Y} au "
                f"{(premier - timedelta(days=1)):%d/%m/%Y} : l'export historique s'arrete "
                f"le {hist_fin:%d/%m/%Y}, le premier rapport mensuel commence le "
                f"{premier:%d/%m/%Y}. Les totaux d'une periode qui chevauche ce trou sont "
                "donc incomplets.")
        dernier = max(i["fin"] for i in infos_lots)
        _alertes.append(
            f"La journee du {dernier:%d/%m/%Y} est incomplete : sa fin de nuit (apres "
            "minuit) figure dans le rapport du mois suivant, qui la completera.")
        _alertes.append(
            f"Reglements, caissiers, profil horaire de reference, inventaire et totaux POS "
            f"restent ceux de l'export historique, arrete au {hist_fin:%d/%m/%Y} : les "
            "rapports mensuels ne les fournissent pas sous la meme forme.")
        _alertes.append(
            "Pour les mois issus d'un rapport mensuel, les quantites vendues et offertes du "
            "compte de resultat se separent par le CA de la ligne (nul = offert). Sur "
            "l'historique, cette regle retrouve la balance a 0,1 % pres sur le total.")

    for r in reconciliation:
        print(f"      [{r['statut']:<16}] {r['controle']}")

    # =================================================================
    # CONSTRUCTION DU JSON
    # =================================================================
    print("\nConstruction du JSON...")
    resultat_jour = [
        {"d": jstr(r["jour"]), "a": r["activite"], "cb": round(r["ca_brut"]),
         "rm": round(r["remise"]), "cn": round(r["ca_net"]), "qv": round(r["qv"]),
         "qo": round(r["qo"]), "ae": round(r["achats"]), "op": round(r["opex"]),
         "cx": round(r["capex"]), "mb": round(r["mb"]), "re": round(r["re"]),
         "rn": round(r["rn"])}
        for r in res]
    ventes_categorie_jour = [
        {"d": jstr(g["jour"]), "t": g["Type"], "c": g["Categorie"], "q": round(g["Qté"]),
         "cn": round(g["CAF"]), "cb": round(g["ca_brut_ligne"]), "rm": round(g["Remise"])}
        for g in regrouper(v, ["jour", "Type", "Categorie"],
                           ["Qté", "CAF", "ca_brut_ligne", "Remise"])]
    articles = [
        {"t": a["Type"], "c": a["Categorie"], "a": a["Articles"], "q": round(a["Qté"]),
         "cn": round(a["CAF"]), "cb": round(a["ca_brut_ligne"]), "rm": round(a["Remise"]),
         "nl": int(a["nl"])}
        for a in art]
    articles_mois = [
        {"m": g["mois"], "t": g["Type"], "c": g["Categorie"], "a": g["Articles"],
         "q": round(g["Qté"]), "cn": round(g["CAF"]), "cb": round(g["ca_brut_ligne"]),
         "rm": round(g["Remise"])}
        for g in regrouper(v, ["mois", "Type", "Categorie", "Articles"],
                           ["Qté", "CAF", "ca_brut_ligne", "Remise"])]
    horaire_jour = [
        {"d": jstr(g["jour"]), "h": int(g["h"]), "on": int(g["ordre_nuit"]), "t": g["Type"],
         "q": round(g["Qté"]), "cn": round(g["CAF"])}
        for g in regrouper(v, ["jour", "h", "ordre_nuit", "Type"], ["Qté", "CAF"])]
    depenses_jour = [
        {"d": jstr(g["jour"]), "n": g["nature"], "p": g["poste"], "m": round(g["montant"])}
        for g in regrouper(dep, ["jour", "nature", "poste"], ["montant"])]
    offerts_jour = [
        {"d": jstr(g["jour"]), "c": g["Caissier"], "cat": g["Catégorie"],
         "q": round(g["Quantité"]), "v": round(g["Val vente"])}
        for g in regrouper(off, ["jour", "Caissier", "Catégorie"], ["Quantité", "Val vente"])]
    oa = regrouper(off, ["Catégorie", "Article"], ["Quantité", "Val vente"])
    oa.sort(key=lambda g: -g["Val vente"])
    offerts_article = [{"cat": g["Catégorie"], "a": g["Article"], "q": round(g["Quantité"]),
                        "v": round(g["Val vente"])} for g in oa[:150]]

    return {
        "meta": {
            "genere_le": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "periode_debut": jstr(jours_v[0]),
            "periode_fin": jstr(jours_v[-1]),
            "jours_exploitation": len(jours_v),
            "donnees_figees": True,
            "source": ("Exports POS TABOO du 26/08/2026 — jeu de donnees fige. "
                       "Regles metier : voir 20_extraction_v2.py."
                       if not (infos_lots or F_CLASSEUR) else
                       "Exports POS TABOO du 26/08/2026"
                       + "".join(f" + rapport mensuel {i['fin']:%m-%Y}" for i in infos_lots)
                       + (f" ; livre de depenses : {F_CLASSEUR.name}" if F_CLASSEUR else "")
                       + ". Regles metier : voir 21_extraction_taboo.py."),
            **({"cumuls_au": jstr(hist_fin)} if infos_lots else {}),
        },
        "totaux_pos": {
            "nb_panier": nb_panier, "panier_moyen": panier_moyen,
            "nombre_vendu": nombre_vendu, "remise_pos": remise_pos,
            "total_offert": total_offert, "ca_ht_hors_offert": ca_ht_hors_offert,
            "ca_offert": ca_offert,
        },
        "pareto": {"n_articles_80pct": n80, "n_articles_total": len(art)},
        "inventaire_meta": {"comptage_saisi": inventaire_compte, "n_sous_alerte": n_alerte,
                            "valeur_stock": round(val_stock)},
        "resultat_jour": resultat_jour,
        "ventes_categorie_jour": ventes_categorie_jour,
        "articles": articles,
        "articles_mois": articles_mois,
        "horaire_jour": horaire_jour,
        "depenses_jour": depenses_jour,
        "offerts_jour": offerts_jour,
        "offerts_article": offerts_article,
        "couts_articles": couts_articles,
        "inventaire": inventaire,
        "reglements": regl,
        "panier_horaire": panier_horaire,
        "caissiers": caissiers,
        "reconciliation": reconciliation,
        "alertes_qualite": _alertes,
    }


def nettoyer_nan(o):
    """NaN/Inf -> None : json.dump ecrirait un NaN que JSON.parse() rejette."""
    if isinstance(o, float) and (o != o or o in (float("inf"), float("-inf"))):
        return None
    if isinstance(o, dict):
        return {k: nettoyer_nan(x) for k, x in o.items()}
    if isinstance(o, list):
        return [nettoyer_nan(x) for x in o]
    return o


def main():
    ap = argparse.ArgumentParser(description="Extraction TABOO sans pandas.")
    ap.add_argument("--sortie", type=Path, default=BASE / "data_v2.json")
    ap.add_argument("--comme-v2", action="store_true",
                    help="ignore rapports mensuels et classeur : doit redonner exactement "
                         "la sortie de 20_extraction_v2.py (test de non-regression)")
    args = ap.parse_args()

    data = nettoyer_nan(extraire(comme_v2=args.comme_v2))
    with open(args.sortie, "w", encoding="utf-8") as f:
        # allow_nan=False : echoue franchement plutot que d'ecrire un JSON invalide.
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"), allow_nan=False)

    print(f"\n{'=' * 70}")
    print(f"{args.sortie.name} ecrit : {args.sortie.stat().st_size/1024:.0f} Ko")
    for k, x in data.items():
        if isinstance(x, list):
            print(f"  {k:<24} {len(x):>7,} lignes")
    if _alertes:
        print("\nALERTES QUALITE reportees dans le dashboard :")
        for a in _alertes:
            print(f"  - {a}")
    print("=" * 70)


if __name__ == "__main__":
    main()
