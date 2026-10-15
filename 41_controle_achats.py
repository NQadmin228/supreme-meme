#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TABOO — Alimente le classeur du controle de gestion a partir des rapports
mensuels de caisse, et rapproche la caisse de l'analyse mensuelle.

  python 41_controle_achats.py

CE QU'IL ECRIT, DANS TABOO_MP_prix_unitaires_et_recettes.xlsx
-------------------------------------------------------------
  Correspondance_caisse   chaque compte de sortie de caisse -> un article de
                          l'onglet Achats, ou un statut (fournisseur non
                          detaille, boisson, charge, a classer) et un poste
                          du livre. C'est l'onglet a corriger : ses choix
                          l'emportent sur les propositions de ce script.
  Journal_achats          une ligne par sortie de caisse rattachee a un
                          article, avec la reference de la sortie. Rejouable :
                          une sortie deja presente n'est jamais ajoutee deux fois.
  Suivi_MP                mois de depart ramene au premier mois disponible,
                          CA du mois par activite renseigne.
  Rapprochement_caisse    la caisse confrontee a l'analyse mensuelle, jour
                          par jour : CA, encaissements, sorties de caisse.

ET POUR LE SITE, taboo-vercel/donnees/taboo/controle_achats.json
  Le journal tel que le classeur le tient (quantites saisies comprises),
  les sorties de caisse par statut et par poste, le rapprochement jour par
  jour. La page Achats matiere le lit ; 60_publier_donnees.py le publie
  avec les autres jeux. Ordre : 21, 22, 41, puis 60.

POURQUOI LE RAPPROCHEMENT EST ORIENTE
-------------------------------------
L'analyse mensuelle (TABOO_Analyses mensuelles.xlsx) fait foi. La caisse,
elle, peut avoir ete corrigee apres coup. Le classeur ne « corrige » donc
rien : il montre ou la caisse s'ecarte de l'analyse, et de combien.

LIMITE ASSUMEE
--------------
La caisse ne donne ni quantite ni fournisseur : seulement un compte, une
date et un montant. Le journal laisse donc la quantite vide -- la deviner
a partir d'un prix de reference rendrait le controle de prix circulaire.
Et la caisse ne voit que les sorties du tiroir, environ la moitie du livre :
les achats regles par virement ou mobile money n'y figurent pas.
"""

import importlib.util
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from sources import source

BASE = Path(__file__).resolve().parent
# Le jeu publie avec les autres par 60_publier_donnees.py, lu par la page
# Achats matiere. 22_decouper.py le conserve quand il refait le dossier.
JEU = BASE / "taboo-vercel" / "donnees" / "taboo" / "controle_achats.json"
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# Les lectures du rapport mensuel et du classeur vivent dans l'extraction :
# une seule facon de lire un fichier, quel que soit l'usage.
_spec = importlib.util.spec_from_file_location("extraction", BASE / "21_extraction_taboo.py")
X = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(X)

ARTICLE, FOURNISSEUR, GENERIQUE, BOISSON, CHARGE, A_CLASSER = (
    "Article", "Fournisseur (non détaillé)", "Générique (non attribuable)",
    "Boisson", "Charge", "À classer")

# Propositions initiales : compte de caisse -> (statut, article, poste du livre,
# hypothese). « Hypothese » signale un rattachement plausible mais non certain,
# a confirmer par le controle de gestion dans l'onglet Correspondance_caisse.
PROPOSITIONS = {
    # --- Matieres du zoom
    "HUILE 25L": (ARTICLE, "Huile", "EAT", False),
    "RIZ": (ARTICLE, "Riz", "EAT", False),
    "FILET DE BOEUF": (ARTICLE, "Filet de bœuf", "EAT", False),
    "GIGOT": (ARTICLE, "Gigot", "EAT", False),
    "BOUTEILLE DE GAZ": (ARTICLE, "Gaz (bouteille)", "GAZ", False),
    "BOX ALU": (ARTICLE, "Emballages alu", "EAT", False),
    "TAKE AWAY": (ARTICLE, "Emballages papier", "EAT", True),
    "SUCRE": (ARTICLE, "Sucre à café", "EAT", True),
    "MANGUE": (ARTICLE, "Mangue (bar)", "DRINK", True),
    "CHARBONS": (ARTICLE, "Charbon", "SMOKE", False),
    "SHISHA": (ARTICLE, "Arôme chicha", "SMOKE", True),
    # --- Autres matieres de la nomenclature
    "PILON": (ARTICLE, "Pilon de poulet", "EAT", False),
    "CUISSE": (ARTICLE, "Cuisse de poulet", "EAT", False),
    "CUISSE ENTIER": (ARTICLE, "Cuisse de poulet", "EAT", True),
    "AILE DE POULET": (ARTICLE, "Ailes de poulet", "EAT", False),
    "WINGS": (ARTICLE, "Ailes de poulet", "EAT", True),
    "MAYONNAISE": (ARTICLE, "Mayonnaise", "EAT", False),
    "CREVETTES": (ARTICLE, "Crevette", "EAT", False),
    "STEAK": (ARTICLE, "Steak burger", "EAT", True),
    "STEAK TABOO": (ARTICLE, "Steak burger", "EAT", True),
    "SPAGHETTI": (ARTICLE, "Spaghetti", "EAT", False),
    "PAIN": (ARTICLE, "Pain baguette", "EAT", True),
    "OIGNON": (ARTICLE, "Oignon", "EAT", False),
    "TOMATES": (ARTICLE, "Tomate", "EAT", False),
    "AVOCAT": (ARTICLE, "Avocat", "EAT", False),
    "PIMENT VERT": (ARTICLE, "Piment", "EAT", False),
    "FARINE": (ARTICLE, "Farine", "EAT", False),
    "VINAIGRE": (ARTICLE, "Vinaigre blanc", "EAT", True),
    "SAUCE SAUJA": (ARTICLE, "Sauce soja", "EAT", False),
    "N'DOLE": (ARTICLE, "Feuilles de ndolé", "EAT", False),
    "LAIT PEAK": (ARTICLE, "Lait", "EAT", False),
    "DORADE": (ARTICLE, "Poisson entier", "EAT", True),
    # --- Achats de matiere sans article precis
    "LEGUMES": (GENERIQUE, None, "EAT", False),
    "FRUIT": (GENERIQUE, None, "DRINK", True),
    "JUS": (GENERIQUE, None, "DRINK", True),
    "BOUILLON": (GENERIQUE, None, "EAT", False),
    "PIMENT EN POUDRE": (GENERIQUE, None, "EAT", False),
    "EPICE": (GENERIQUE, None, "EAT", False),
    "AROME VANILLE": (GENERIQUE, None, "EAT", False),
    # --- Fournisseurs : un ticket, plusieurs articles inconnus
    "LEA SHOP": (FOURNISSEUR, None, "EAT", True),
    "CHAMPION": (FOURNISSEUR, None, "EAT", True),
    "DAFRIMEX": (FOURNISSEUR, None, "DRINK", False),
    "WARE HOUSE": (FOURNISSEUR, None, "DRINK", True),
    "BOUBA SHOP": (FOURNISSEUR, None, "EAT", True),
    "RAMCO": (FOURNISSEUR, None, "EAT", True),
    "BENEDICTE": (FOURNISSEUR, None, "EAT", True),
    "CALIENDI": (FOURNISSEUR, None, "DRINK", True),
    "PAPYRUS": (FOURNISSEUR, None, "CONSOMMABLES", True),
    "ETS BAIMEX": (FOURNISSEUR, None, "EAT", True),
    "JEAN AZAR": (FOURNISSEUR, None, "EAT", True),
    # --- Boissons et glace
    **{b: (BOISSON, None, "DRINK", False) for b in (
        "DESPERADOS", "RED BULL", "HEINEKEN", "COCA", "CRISTAL", "MARTELL", "GUINNESS",
        "JACK DANIEL", "BEAUFORT", "JAGERMEISTER", "SPRITE", "FUTS DE BIERRE",
        "COCKTAIL", "PURE WATER")},
    "GLACON": (BOISSON, None, "GLACONS", False),
    # --- Charges : rien a voir avec la matiere premiere
    **{c: (CHARGE, None, "ENTRETIEN", False) for c in (
        "VIDANGE", "ENTRETIEN GENERALE", "CURAGE", "FRIGORISTE", "ETANCHEISTE",
        "ENTRETIEN", "PLOMBIER", "MAIN D'0EUVRE PLOMBIER", "DEBOUCHAGE WC",
        "INSECTICIDES", "ACHAT DE MATERIEL ENTRETIEN", "DETERGENT", "JAVEL",
        "LIQUIDE VAISSELLE", "ORDURES", "MAINTENANCE ANTENNE", "ESSENCE CITRONNELLE")},
    **{c: (CHARGE, None, "TRAVAUX", True) for c in (
        "MENUSIER DECORATEUR", "MENUSIER ALU", "REVETEMENT TASSO", "CIMENT", "SABLE")},
    **{c: (CHARGE, None, "EQUIPEMENTS", True) for c in (
        "RECHAUD", "CHAUFFE CHARBON", "EXTENDER", "RAPEUSE", "AGRAPHEUSE", "PILE MICROS")},
    **{c: (CHARGE, None, "CONSOMMABLES", True) for c in (
        "PAPIER THERMIQUE", "GANTS", "RED CUP", "PAILLES", "CACHE NEZ", "SACHET PORTION",
        "SACHET BLANC", "TIGE BROCHETTES", "PHOTOCOPIE")},
    **{c: (CHARGE, None, "TRANSPORT", False) for c in (
        "CARBURANT", "LIVRAISON", "DEPLACEMENT", "DEPLACEMENT JUSTIN",
        "DEPLACEMENT ELOM", "DEPLACEMENT ESAIE", "DEPLACEMENT BIERRE")},
    **{c: (CHARGE, None, "INTERNET / TV", False) for c in (
        "ABONNEMENT WIFI", "ABONNEMENT CANAL", "CANAL PLUS CESAR")},
    "CACHET DJ": (CHARGE, None, "CACHETS", False),
    "MONNAIE": (CHARGE, None, "MONNAIE", False),
    "CARTE": (CHARGE, None, "TELEPHONIE", True),
    "FRAIS DE TRANSACTIN": (CHARGE, None, "ADMINISTRATIF", True),
}

POS_VERS_CLASSEUR = [("Espèces", "ESPECES", "CASH"), ("Carte", "CB", "TPE"),
                     ("Mixx (TMONEY)", "TMONEY", "MIXX"), ("Flooz", "FLOOZ", "FLOOZ"),
                     ("Gozem", "GOZEM", "GOZEM")]
POSTES_ACTIVITE = {"EAT": ("EAT", "GAZ"), "DRINK": ("DRINK", "MIAMI 228", "PICASSO", "GLACONS"),
                   "SMOKE": ("SMOKE",)}


def jours_du_lot(lot):
    """Les dates calendaires couvertes par un rapport mensuel (journal de vente)."""
    return sorted({d.date() for d in (X.date_heure(r["Date"], "%d/%m/%Y %H:%M:%S")
                                      for r in lot["ventes"]) if d})


# -----------------------------------------------------------------
# Enregistrer sans rien perdre de ce qu'Excel a ajoute
# -----------------------------------------------------------------
# Une validation ou une mise en forme conditionnelle qui pointe vers un AUTRE
# onglet (liste d'articles = Achats!$B$6:$B$134, seuils = Parametres!$B$10),
# Excel l'enregistre dans une extension x14, en fin de feuille. openpyxl ne
# relit pas ces extensions : un simple load/save les efface sans un mot, et
# avec elles les listes deroulantes d'articles et les alertes de Fiches_cout.
# On les reporte donc telles quelles du fichier d'origine vers le fichier
# ecrit, puis on recompte : s'il manque quoi que ce soit, le classeur n'est
# pas remplace.
NS_XL = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
         "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
         "x14": "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main",
         "pr": "http://schemas.openxmlformats.org/package/2006/relationships"}
EXT_RECONDUITES = {"{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}": "validations x14",
                   "{78C0D931-6437-407d-A8EE-F0AAD7539E65}": "mises en forme conditionnelles x14"}


def _parties_onglets(z):
    """Nom d'onglet -> chemin de sa feuille dans le paquet xlsx."""
    import posixpath
    import xml.etree.ElementTree as ET
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    cible = {r.get("Id"): r.get("Target") for r in rels.findall("pr:Relationship", NS_XL)}
    parties = {}
    for s in wb.find("m:sheets", NS_XL):
        t = cible[s.get(f"{{{NS_XL['r']}}}id")]
        parties[s.get("name")] = t.lstrip("/") if t.startswith("/") else posixpath.join("xl", t)
    return parties


def inventaire_classeur(chemin):
    """Par onglet : validations et regles de mise en forme, standard et x14 ;
    et les types de contenu du paquet (graphiques, images, tableaux...)."""
    import xml.etree.ElementTree as ET
    import zipfile
    onglets = {}
    with zipfile.ZipFile(chemin) as z:
        for nom, partie in _parties_onglets(z).items():
            x = ET.fromstring(z.read(partie))
            onglets[nom] = {
                "validations": len(x.findall("m:dataValidations/m:dataValidation", NS_XL)),
                "validations x14": len(x.findall(
                    "m:extLst/m:ext/x14:dataValidations/x14:dataValidation", NS_XL)),
                "mises en forme": len(x.findall("m:conditionalFormatting/m:cfRule", NS_XL)),
                "mises en forme x14": len(x.findall(
                    "m:extLst/m:ext/x14:conditionalFormattings/x14:conditionalFormatting/x14:cfRule",
                    NS_XL)),
                "extensions": [e.get("uri") for e in x.findall("m:extLst/m:ext", NS_XL)]}
        types = {o.get("ContentType") for o in ET.fromstring(z.read("[Content_Types].xml"))
                 if o.get("ContentType")}
    return onglets, types


def enregistrer_sans_perte(wb, chemin, recrees):
    """Enregistre wb a la place de chemin sans perdre les extensions x14.

    Les onglets de `recrees` sont reecrits de zero par ce script : on n'y reporte
    rien, mais leurs validations et mises en forme ne doivent pas etre moins
    nombreuses qu'avant (Excel les avait seulement passees en x14). Pour les
    autres, chaque compte doit etre au moins celui d'avant."""
    import os
    import zipfile

    avant, types_avant = inventaire_classeur(chemin)
    inconnues = [(o, u) for o, i in avant.items() if o not in recrees
                 for u in i["extensions"] if u not in EXT_RECONDUITES]
    if inconnues:
        raise SystemExit(
            "Classeur non modifié : il contient des extensions Excel qu'openpyxl effacerait "
            "et que ce script ne sait pas reporter : "
            + ", ".join(f"{o} {u}" for o, u in inconnues)
            + ". Si elles se reportent telles quelles, ajouter leur uri à EXT_RECONDUITES.")
    # Chaque extension est reprise octet pour octet ; seuls s'y ajoutent les
    # espaces de noms qu'elle emprunte a la racine de sa feuille (xr:uid...).
    extensions = {}
    with zipfile.ZipFile(chemin) as z:
        for nom, partie in _parties_onglets(z).items():
            if nom in recrees:
                continue
            xml = z.read(partie).decode("utf-8")
            if "<extLst>" not in xml:
                continue
            racine = dict(re.findall(r'xmlns:(\w+)="([^"]+)"', re.search(r"<worksheet\b[^>]*>", xml)[0]))
            fin = xml[xml.rfind("<extLst>"):]
            extensions[nom] = []
            for m in re.finditer(r'<ext\b[^>]*\buri="([^"]+)"[^>]*>.*?</ext>', fin, re.S):
                if m[1] not in EXT_RECONDUITES:
                    continue
                bloc = m[0]
                manquants = (set(re.findall(r'</?(\w+):\w', bloc)) | set(re.findall(r'\s(\w+):\w+=', bloc))) \
                    - set(re.findall(r'xmlns:(\w+)=', bloc)) - {"xmlns"}
                if manquants - set(racine):
                    raise SystemExit(f"Classeur non modifié : extension {m[1]} de {nom} illisible "
                                     f"(préfixes {sorted(manquants - set(racine))}).")
                decl = "".join(f' xmlns:{p}="{racine[p]}"' for p in sorted(manquants))
                extensions[nom].append((m[1], bloc.replace("<ext ", f"<ext{decl} ", 1)))

    brut = chemin.with_name(chemin.stem + ".openpyxl.tmp")
    final = chemin.with_name(chemin.stem + ".en_cours.tmp")
    try:
        wb.save(brut)
        with zipfile.ZipFile(brut) as zs, zipfile.ZipFile(final, "w", zipfile.ZIP_DEFLATED) as zd:
            parties = {p: n for n, p in _parties_onglets(zs).items()}
            for info in zs.infolist():
                donnees = zs.read(info)
                nom = parties.get(info.filename)
                a_reporter = extensions.get(nom) if nom else None
                if a_reporter:
                    xml = donnees.decode("utf-8")
                    deja = set(re.findall(r'<ext [^>]*uri="([^"]+)"', xml))
                    bloc = "".join(b for u, b in a_reporter if u not in deja)
                    if "</extLst></worksheet>" in xml:
                        xml = xml.replace("</extLst></worksheet>", bloc + "</extLst></worksheet>")
                    else:
                        xml = xml.replace("</worksheet>", f"<extLst>{bloc}</extLst></worksheet>")
                    donnees = xml.encode("utf-8")
                zd.writestr(info, donnees)

        apres, types_apres = inventaire_classeur(final)
        pertes = []
        for nom, i in avant.items():
            j = apres.get(nom)
            if j is None:
                pertes.append(f"onglet {nom}")
                continue
            if nom in recrees:
                for genre in ("validations", "mises en forme"):
                    if j[genre] + j[genre + " x14"] < i[genre] + i[genre + " x14"]:
                        pertes.append(f"{nom} : {genre} {i[genre] + i[genre + ' x14']} -> "
                                      f"{j[genre] + j[genre + ' x14']}")
            else:
                for genre in ("validations", "validations x14", "mises en forme",
                              "mises en forme x14"):
                    if j[genre] < i[genre]:
                        pertes.append(f"{nom} : {genre} {i[genre]} -> {j[genre]}")
        # La chaine de calcul, Excel la reconstruit, et openpyxl ecrit les textes
        # dans les cellules plutot qu'en table partagee ; tout autre type de
        # contenu disparu (graphique, image, tableau...) est une perte.
        pertes += [f"partie {t}" for t in sorted(types_avant - types_apres)
                   if not t.endswith(("calcChain+xml", "sharedStrings+xml"))]
        if pertes:
            raise SystemExit("Classeur non modifié : l'enregistrement perdrait "
                             + " ; ".join(pertes))
        os.replace(final, chemin)
    finally:
        for f in (brut, final):
            if f.exists():
                f.unlink()
    return apres


def main():
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    lots = X.lots_mensuels()
    if not lots:
        raise SystemExit("Aucun rapport mensuel configure (TABOO_RAPPORTS_MENSUELS).")
    F_ANALYSE = source("TABOO_ANALYSE_MENSUELLE")
    F_MP = source("TABOO_CLASSEUR_MP")
    data = json.loads((BASE / "data_v2.json").read_text(encoding="utf-8"))

    wb = openpyxl.load_workbook(F_MP)
    police = wb["Journal_achats"]["A5"].font.name or "Arial"
    F_TITRE = Font(name=police, size=13, bold=True, color="FF1B2A49")
    F_NOTE = Font(name=police, size=9, italic=True, color="FF595959")
    F_ENTETE = Font(name=police, bold=True, color="FFFFFFFF")
    F_SAISIE = Font(name=police, color="FF0000FF")
    F_CALC = Font(name=police, color="FF000000")
    F_GRAS = Font(name=police, bold=True)
    FOND_ENTETE = PatternFill("solid", fgColor="FF1B2A49")
    FOND_JAUNE = PatternFill("solid", fgColor="FFFFFF99")
    FOND_GRIS = PatternFill("solid", fgColor="FFF2F2F2")
    FOND_ALERTE = PatternFill("solid", fgColor="FFFCE4D6")
    FMT_F = "#,##0;\\(#,##0\\);\\-"
    FMT_P = "0.0%"
    ENROULE = Alignment(wrap_text=True, vertical="top")

    articles = {wb["Achats"].cell(r, 2).value for r in range(6, 251)} - {None}

    def entete(ws, ligne, titres, largeurs=None):
        for i, t in enumerate(titres, start=1):
            c = ws.cell(ligne, i, t)
            c.font, c.fill, c.alignment = F_ENTETE, FOND_ENTETE, Alignment(wrap_text=True, vertical="center")
            if largeurs:
                ws.column_dimensions[get_column_letter(i)].width = largeurs[i - 1]

    # -----------------------------------------------------------------
    # 1. Correspondance des comptes de caisse
    # -----------------------------------------------------------------
    depenses = []
    for lot in lots:
        for r in lot["depenses"]:
            d = X.date_heure(r["Date"], "%d/%m/%Y %H:%M:%S")
            if d is None:
                continue
            depenses.append({"id": r["ID"], "date": d, "compte": (r["Compte"] or "").strip(),
                             "montant": X.z(X.num(r["Montant"])), "caissier": r["Caissier"],
                             "commentaire": r["Commentaire"]})
    nb, tot = Counter(), Counter()
    for r in depenses:
        nb[r["compte"]] += 1
        tot[r["compte"]] += r["montant"]

    NOM_CORR = "Correspondance_caisse"
    choix = {}
    if NOM_CORR in wb.sheetnames:
        ws = wb[NOM_CORR]
        for r in range(6, ws.max_row + 1):
            compte = ws.cell(r, 1).value
            if compte:
                choix[str(compte).strip()] = (ws.cell(r, 4).value, ws.cell(r, 5).value,
                                             ws.cell(r, 6).value, ws.cell(r, 7).value == "Oui")
        del wb[NOM_CORR]
    for compte in nb:
        if compte not in choix:
            choix[compte] = PROPOSITIONS.get(compte, (A_CLASSER, None, None, False))
    for compte, (statut, article, poste, hyp) in choix.items():
        if statut == ARTICLE and article not in articles:
            raise SystemExit(f"Correspondance : « {article} » (compte {compte}) n'est pas "
                             "un article de l'onglet Achats.")

    ws = wb.create_sheet(NOM_CORR, index=wb.sheetnames.index("Journal_achats"))
    ws["A1"] = "Correspondance des sorties de caisse – compte → article et poste du livre"
    ws["A1"].font = F_TITRE
    ws["A2"] = ("La caisse ne nomme qu'un « compte » : parfois un article (RIZ), parfois un "
                "fournisseur (LEA SHOP), parfois une charge (VIDANGE). Cet onglet dit ce que "
                "chacun représente. Les colonnes en bleu se corrigent ici ; 41_controle_achats.py "
                "relit ces choix à chaque passage et ne les écrase jamais.")
    ws["A2"].font, ws["A2"].alignment = F_NOTE, ENROULE
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 42
    ws["A3"] = "« Hypothèse = Oui » : rattachement plausible, à confirmer."
    ws["A3"].font = F_NOTE
    entete(ws, 5, ["Compte de caisse", "Nombre de sorties", "Montant (FCFA)", "Statut",
                   "Article (onglet Achats)", "Poste du livre", "Hypothèse", "Dans le journal"],
           [30, 11, 14, 26, 26, 18, 11, 14])
    ordre = sorted(choix, key=lambda c: (-tot.get(c, 0), c))
    for i, compte in enumerate(ordre, start=6):
        statut, article, poste, hyp = choix[compte]
        vals = [compte, nb.get(compte, 0), tot.get(compte, 0), statut, article, poste,
                "Oui" if hyp else "Non",
                f'=IF(D{i}="{ARTICLE}","Oui","Non")']
        for j, x in enumerate(vals, start=1):
            c = ws.cell(i, j, x)
            c.font = F_SAISIE if j in (4, 5, 6, 7) else F_CALC
            if j == 3:
                c.number_format = FMT_F
            if statut == A_CLASSER and j == 4:
                c.fill = FOND_JAUNE
            if hyp and j == 7:
                c.fill = FOND_JAUNE
    fin = 5 + len(ordre)
    ws.freeze_panes = "B6"
    ws.auto_filter.ref = f"A5:H{fin}"
    dv = DataValidation(type="list", formula1='"' + ",".join(
        [ARTICLE, FOURNISSEUR, GENERIQUE, BOISSON, CHARGE, A_CLASSER]) + '"')
    dv_art = DataValidation(type="list", formula1="=Achats!$B$6:$B$134", allow_blank=True)
    dv_oui = DataValidation(type="list", formula1='"Oui,Non"')
    for v_, plage in ((dv, f"D6:D{fin + 50}"), (dv_art, f"E6:E{fin + 50}"),
                      (dv_oui, f"G6:G{fin + 50}")):
        ws.add_data_validation(v_)
        v_.add(plage)

    # -----------------------------------------------------------------
    # 2. Journal des achats
    # -----------------------------------------------------------------
    wj = wb["Journal_achats"]
    deja = {str(wj.cell(r, 3).value) for r in range(6, 506) if wj.cell(r, 3).value}
    libre = next(r for r in range(6, 507) if r == 506 or wj.cell(r, 1).value is None)
    ajoutees = 0
    for r in sorted(depenses, key=lambda x: (x["date"], int(x["id"]))):
        statut, article, _, _ = choix[r["compte"]]
        if statut != ARTICLE:
            continue
        ref = f"Caisse #{r['id']} – {r['caissier']}"
        if ref in deja:
            continue
        if libre > 505:
            raise SystemExit("Journal_achats est plein (500 lignes) : l'agrandir avant de rejouer.")
        for col, x, fmt in ((1, r["date"], "dd/mm/yyyy"), (2, article, None), (3, ref, None),
                            (5, r["montant"], FMT_F)):
            c = wj.cell(libre, col, x)
            c.font = F_SAISIE
            if fmt:
                c.number_format = fmt
        ajoutees += 1
        libre += 1
    wj["C5"].comment = Comment(
        "Lignes « Caisse #ID – caissier » : reprises des sorties de caisse du rapport mensuel "
        "(états des dépenses) par 41_controle_achats.py. La caisse ne donne ni fournisseur ni "
        "quantité : la quantité reste à saisir pour obtenir le prix unitaire payé.",
        "41_controle_achats.py")

    # -----------------------------------------------------------------
    # 3. Suivi_MP : mois de depart et CA du mois
    # -----------------------------------------------------------------
    wm = wb["Suivi_MP"]
    premier = min(jours_du_lot(l)[0] for l in lots)
    depart = datetime(premier.year, premier.month, 1)
    if not isinstance(wm["B3"].value, datetime) or wm["B3"].value > depart:
        wm["B3"] = depart
    depart = wm["B3"].value
    ca = defaultdict(float)
    for r in data["resultat_jour"]:
        ca[(r["d"][:7], r["a"])] += r["cn"]
    for ligne, act in ((58, "EAT"), (61, "DRINK"), (64, "SMOKE")):
        for k in range(12):
            an, mo = divmod(depart.month - 1 + k, 12)
            cle_mois = f"{depart.year + an}-{mo + 1:02d}"
            if cle_mois > data["meta"]["periode_fin"][:7]:
                break
            if cle_mois < data["meta"].get("cumuls_au", "0000")[:7] or (cle_mois, act) not in ca:
                continue
            c = wm.cell(ligne, 3 + k, round(ca[(cle_mois, act)]))
            c.font, c.number_format = F_SAISIE, FMT_F
    wm["B58"].comment = Comment(
        "CA net du mois par activité, repris du dashboard TABOO (journal de vente de la "
        "caisse) par 41_controle_achats.py. L'analyse mensuelle n'a pas le détail par "
        "activité ; son total peut différer des remises accordées sur le ticket entier "
        "(voir Rapprochement_caisse).", "41_controle_achats.py")

    # -----------------------------------------------------------------
    # 4. Rapprochement caisse / analyse mensuelle
    # -----------------------------------------------------------------
    NOM_RAP = "Rapprochement_caisse"
    if NOM_RAP in wb.sheetnames:
        del wb[NOM_RAP]
    wr = wb.create_sheet(NOM_RAP)
    _, livre = X.lire_livre_classeur(F_ANALYSE)
    livre_jour = defaultdict(lambda: defaultdict(float))
    for r in livre:
        d = X.date_heure(r.get("DATES"), "%d/%m/%Y")
        if d is None or (r.get("JOURS") or "").lower() not in X.JOURS_SEMAINE:
            continue
        for poste in list(X.MAPPING_ACHATS) + X.POSTES_OPEX + X.POSTES_CAPEX:
            m = X.num(r.get(poste))
            if m:
                livre_jour[d.date()][poste] += m

    rappro_jours = []
    ligne = 1
    wr.cell(ligne, 1, "Rapprochement : la caisse face à l'analyse mensuelle").font = F_TITRE
    ligne += 1
    c = wr.cell(ligne, 1, "L'analyse mensuelle (TABOO_Analyses mensuelles.xlsx) fait foi ; la "
                          "caisse peut avoir été corrigée après coup. Ce tableau ne corrige rien : il "
                          "montre où la caisse s'écarte de l'analyse, et de combien. Valeurs en bleu = "
                          "reprises des fichiers sources ; colonnes noires = calculées. Les espèces ne "
                          "se comparent pas jour par jour : l'analyse déduit le CASH du CA (CA moins les "
                          "autres moyens et les impayés), la caisse date chaque paiement au jour où il est "
                          "réglé. Le contrôle du cash est celui de l'analyse elle-même, son vérificateur "
                          "(différentiel cash), qui doit valoir 0.")
    c.font, c.alignment = F_NOTE, ENROULE
    wr.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=20)
    wr.row_dimensions[ligne].height = 45
    ligne += 2
    for lot in lots:
        jours = jours_du_lot(lot)
        fin_mois = jours[-1]
        onglet, ref = X.onglet_du_mois(F_ANALYSE, fin_mois.year, fin_mois.month)
        mois = f"{X.MOIS_FR[fin_mois.month - 1].capitalize()} {fin_mois.year}"
        ca_caisse = defaultdict(float)
        for r in lot["ventes"]:
            d = X.date_heure(r["Date"], "%d/%m/%Y %H:%M:%S")
            if d and r["Type"] != "TOTAL":
                ca_caisse[d.date()] += X.z(X.num(r["CAF"]))
        caf = {X.date_heure(r["Date"], "%d/%m/%Y %H:%M:%S").date(): r for r in lot["caf"]}
        regl = defaultdict(lambda: defaultdict(float))
        for r in lot["reglements"]:
            d = X.date_heure(r["Date du payement"], "%d/%m/%Y %H:%M:%S")
            if d:
                regl[d.date()][r["Moyen de paiement"]] += X.z(X.num(r["Montant payé"]))
        sorties = defaultdict(float)
        sorties_poste = defaultdict(float)
        for r in depenses:
            if r["date"].date() in set(jours):
                sorties[r["date"].date()] += r["montant"]
                sorties_poste[choix[r["compte"]][2] or "Non classé"] += r["montant"]

        wr.cell(ligne, 1, f"{mois} — jour par jour"
                + (f" (analyse : onglet {onglet.strip()})" if onglet else
                   " — ONGLET ABSENT DE L'ANALYSE MENSUELLE")).font = F_GRAS
        ligne += 1
        titres = ["Date", "CA caisse (journal de vente)", "CA journal CAF avant remise ticket",
                  "Remises sur ticket (journal CAF)", "CA analyse mensuelle",
                  "Écart caisse − analyse", "dont remises sur ticket", "dont écart entre journaux",
                  "Reste inexpliqué", "Espèces caisse", "Espèces analyse (CASH)",
                  "Écart espèces (décalage, pas une alerte)", "Autres moyens caisse",
                  "Autres moyens analyse", "Écart autres moyens (décalage, pas une alerte)",
                  "Sorties de caisse", "Dépenses du livre", "Part passée par la caisse",
                  "Alerte", "Vérificateur de l'analyse (différentiel cash)"]
        entete(wr, ligne, titres, [11, 14, 14, 13, 14, 13, 13, 13, 12, 13, 13, 14, 13, 13, 14,
                                   13, 13, 11, 22, 14])
        wr.row_dimensions[ligne].height = 45
        haut = ligne + 1
        for d in jours:
            ligne += 1
            rf = ref.get(d, {})
            cf = caf.get(d, {})
            autres_c = sum(regl[d].get(m, 0) for _, m, _ in POS_VERS_CLASSEUR[1:])
            autres_a = sum(rf.get(col, 0) for _, _, col in POS_VERS_CLASSEUR[1:])
            vals = {1: datetime(d.year, d.month, d.day), 2: ca_caisse[d],
                    3: X.z(X.num(cf.get("HT"))), 4: X.z(X.num(cf.get("Remise"))),
                    5: rf.get("CA", 0), 10: regl[d].get("ESPECES", 0), 11: rf.get("CASH", 0),
                    13: autres_c, 14: autres_a, 16: sorties[d],
                    17: sum(livre_jour[d].values()), 20: rf.get("DIFERENTIEL CASH", 0)}
            for col, x in vals.items():
                c = wr.cell(ligne, col, x)
                c.font = F_SAISIE
                c.number_format = "dd/mm/yyyy" if col == 1 else FMT_F
            rappro_jours.append({"d": d.isoformat(), **{k: round(vals[col]) for k, col in (
                ("cac", 2), ("caj", 3), ("rem", 4), ("can", 5), ("esc", 10), ("esa", 11),
                ("auc", 13), ("aua", 14), ("sc", 16), ("dl", 17), ("ver", 20))}})
            L = ligne
            formules = {6: f"=B{L}-E{L}", 7: f"=D{L}", 8: f"=B{L}-C{L}", 9: f"=F{L}-G{L}-H{L}",
                        12: f"=J{L}-K{L}", 15: f"=M{L}-N{L}",
                        18: f'=IF(Q{L}>0,P{L}/Q{L},"")',
                        19: (f'=TRIM(IF(ABS(I{L})>=1,"CA inexpliqué ","")'
                             f'&IF(ABS(T{L})>=1,"vérificateur analyse ","")'
                             f'&IF(P{L}>Q{L}+0.5,"caisse > livre",""))')}
            for col, f in formules.items():
                c = wr.cell(ligne, col, f)
                c.font, c.fill = F_CALC, FOND_GRIS
                c.number_format = FMT_P if col == 18 else FMT_F
        bas = ligne
        ligne += 1
        wr.cell(ligne, 1, "Total").font = F_GRAS
        for col in [*range(2, 18), 20]:
            c = wr.cell(ligne, col, f"=SUM({get_column_letter(col)}{haut}:{get_column_letter(col)}{bas})")
            c.font, c.number_format = F_GRAS, FMT_F
        c = wr.cell(ligne, 18, f'=IF(Q{ligne}>0,P{ligne}/Q{ligne},"")')
        c.font, c.number_format = F_GRAS, FMT_P
        c = wr.cell(ligne, 19, f'=COUNTIF(S{haut}:S{bas},"?*")&" jour(s) en alerte"')
        c.font = F_GRAS
        total = ligne
        from openpyxl.formatting.rule import FormulaRule
        wr.conditional_formatting.add(f"A{haut}:T{bas}",
                                      FormulaRule(formula=[f'LEN($S{haut})>0'], fill=FOND_ALERTE))

        # Par moyen de paiement
        ligne += 3
        wr.cell(ligne, 1, f"{mois} — encaissements par moyen de paiement").font = F_GRAS
        ligne += 1
        entete(wr, ligne, ["Moyen", "Caisse", "Analyse mensuelle", "Écart", "Écart en % de l'analyse"])
        for lib, pos, col in POS_VERS_CLASSEUR:
            ligne += 1
            wr.cell(ligne, 1, lib)
            for j, x in ((2, sum(regl[d].get(pos, 0) for d in jours)),
                         (3, sum(ref.get(d, {}).get(col, 0) for d in jours))):
                c = wr.cell(ligne, j, x)
                c.font, c.number_format = F_SAISIE, FMT_F
            c = wr.cell(ligne, 4, f"=B{ligne}-C{ligne}")
            c.font, c.fill, c.number_format = F_CALC, FOND_GRIS, FMT_F
            c = wr.cell(ligne, 5, f'=IF(C{ligne}<>0,D{ligne}/C{ligne},"")')
            c.font, c.fill, c.number_format = F_CALC, FOND_GRIS, FMT_P
        ligne += 1
        wr.cell(ligne, 1, "Impayés (analyse)")
        c = wr.cell(ligne, 3, sum(ref.get(d, {}).get("IMPAYES", 0) for d in jours))
        c.font, c.number_format = F_SAISIE, FMT_F

        # Par poste du livre
        ligne += 3
        wr.cell(ligne, 1, f"{mois} — dépenses par poste : livre et sorties de caisse").font = F_GRAS
        ligne += 1
        entete(wr, ligne, ["Poste du livre", "Livre (analyse)", "Sorties de caisse classées",
                           "Part passée par la caisse", "Lecture"])
        postes_livre = defaultdict(float)
        for d in jours:
            for p, m in livre_jour[d].items():
                postes_livre[p] += m
        for p in sorted(set(postes_livre) | set(sorties_poste), key=lambda p: -postes_livre.get(p, 0)):
            ligne += 1
            wr.cell(ligne, 1, p)
            for j, x in ((2, postes_livre.get(p, 0)), (3, sorties_poste.get(p, 0))):
                c = wr.cell(ligne, j, x)
                c.font, c.number_format = F_SAISIE, FMT_F
            c = wr.cell(ligne, 4, f'=IF(B{ligne}>0,C{ligne}/B{ligne},"")')
            c.font, c.fill, c.number_format = F_CALC, FOND_GRIS, FMT_P
            c = wr.cell(ligne, 5, f'=IF(C{ligne}>B{ligne}+0.5,"caisse > livre : à instruire","")')
            c.font, c.fill = F_CALC, FOND_GRIS

        # Achats du livre rapportes au CA
        ligne += 3
        wr.cell(ligne, 1, f"{mois} — achats du livre rapportés au CA (question 3)").font = F_GRAS
        ligne += 1
        entete(wr, ligne, ["Activité", "Achats du livre", "CA net caisse", "Achats / CA",
                           "Postes du livre retenus"])
        for act, postes in POSTES_ACTIVITE.items():
            ligne += 1
            wr.cell(ligne, 1, act)
            c = wr.cell(ligne, 2, sum(postes_livre.get(p, 0) for p in postes))
            c.font, c.number_format = F_SAISIE, FMT_F
            ca_act = sum(r["cn"] for r in data["resultat_jour"]
                         if r["a"] == act and r["d"][:7] == f"{fin_mois:%Y-%m}")
            c = wr.cell(ligne, 3, ca_act)
            c.font, c.number_format = F_SAISIE, FMT_F
            c = wr.cell(ligne, 4, f'=IF(C{ligne}>0,B{ligne}/C{ligne},"")')
            c.font, c.fill, c.number_format = F_CALC, FOND_GRIS, FMT_P
            wr.cell(ligne, 5, " + ".join(postes)).font = F_NOTE
        ligne += 1
        c = wr.cell(ligne, 1, "Le journal d'achats ne reprend que les sorties de caisse rattachées à "
                              "un article ; ce tableau-ci prend tout le livre, qui fait foi.")
        c.font = F_NOTE
        ligne += 3
    wr.freeze_panes = "B6"

    inv = enregistrer_sans_perte(wb, F_MP, {NOM_CORR, NOM_RAP})

    # -----------------------------------------------------------------
    # 5. Le jeu du site : ce que le classeur sait, une fois corrige
    # -----------------------------------------------------------------
    # Le journal est relu dans le classeur, pas dans la caisse : une
    # quantite saisie par le controle de gestion, une ligne ajoutee a la
    # main, un rattachement corrige dans Correspondance_caisse doivent se
    # voir sur le site. Ni reference de sortie ni nom de caissier : la
    # page n'en a pas besoin.
    compte_de = {str(r["id"]): r["compte"] for r in depenses}
    poste_article = {}
    for statut, article, poste, _ in choix.values():
        if statut == ARTICLE and poste:
            poste_article.setdefault(article, poste)
    wa = wb["Achats"]
    fiches = {}
    for r in range(6, 251):
        nom = wa.cell(r, 2).value
        if not nom:
            continue
        # Les formules H et M de l'onglet, refaites ici : openpyxl n'en lit
        # pas le resultat. Prix du lot = min, ou moyenne min-max ;
        # prix unitaire = prix du lot / quantite du lot / (1 - perte).
        cat, unite, pmin, pmax, lot, perte = (wa.cell(r, c).value for c in (1, 5, 6, 7, 10, 12))
        nombre = lambda x: isinstance(x, (int, float))
        prix = None if not nombre(pmin) else (pmin + pmax) / 2 if nombre(pmax) else pmin
        pu = prix / lot / (1 - (perte or 0)) if prix is not None and nombre(lot) and lot else None
        fiches[nom] = {"u": unite, "cat": cat, "pu": round(pu, 4) if pu else None}
    journal = []
    for r in range(6, 506):
        d, article, ref, q, m = (wj.cell(r, c).value for c in (1, 2, 3, 4, 5))
        if not isinstance(d, datetime) or not article or not isinstance(m, (int, float)):
            continue
        caisse = re.match(r"Caisse #(\d+)", str(ref or ""))
        if caisse and compte_de.get(caisse[1]) in choix:
            poste = choix[compte_de[caisse[1]]][2]
        else:
            poste = poste_article.get(article) or (str(fiches.get(article, {}).get("cat") or "")
                                                   .split(" ")[0] or None)
        journal.append({"d": d.date().isoformat(), "a": article, "p": poste, "m": round(m),
                        "q": q if isinstance(q, (int, float)) else None,
                        "o": "caisse" if caisse else "saisie"})
    sorties = defaultdict(lambda: {"m": 0, "n": 0})
    for r in depenses:
        statut, _, poste, _ = choix[r["compte"]]
        s = sorties[(r["date"].date().isoformat(), statut, poste)]
        s["m"] += r["montant"]
        s["n"] += 1
    jeu = {
        "meta": {"genere_le": f"{datetime.now():%Y-%m-%d %H:%M}",
                 "mois": sorted({f"{jours_du_lot(l)[-1]:%Y-%m}" for l in lots}),
                 "a_classer": sum(1 for c in ordre if choix[c][0] == A_CLASSER),
                 "hypotheses": sum(1 for c in ordre if choix[c][0] == ARTICLE and choix[c][3])},
        "articles": {a: fiches[a] for a in sorted({j["a"] for j in journal}) if a in fiches},
        "journal": journal,
        "sorties": [{"d": d, "s": s, "p": p, "m": round(v["m"]), "n": v["n"]}
                    for (d, s, p), v in sorted(sorties.items(), key=lambda x: (x[0][0], x[0][1], x[0][2] or ""))],
        "rapprochement": rappro_jours,
    }
    JEU.parent.mkdir(parents=True, exist_ok=True)
    JEU.write_text(json.dumps(jeu, ensure_ascii=False, separators=(",", ":"), allow_nan=False),
                   encoding="utf-8")

    print(f"  {F_MP.name}")
    print(f"    Validations / mises en forme x14 reportées : "
          f"{sum(i['validations x14'] for i in inv.values())} / "
          f"{sum(i['mises en forme x14'] for i in inv.values())}")
    print(f"    Correspondance_caisse : {len(ordre)} comptes "
          f"({sum(1 for c in ordre if choix[c][0] == ARTICLE)} rattachés à un article, "
          f"{sum(1 for c in ordre if choix[c][0] == A_CLASSER)} à classer, "
          f"{sum(1 for c in ordre if choix[c][3])} hypothèses)")
    print(f"    Journal_achats        : {ajoutees} lignes ajoutées")
    print(f"    Suivi_MP              : départ {depart:%m/%Y}")
    print(f"    Rapprochement_caisse  : {len(lots)} mois")
    print(f"  {JEU.relative_to(BASE)} : {len(journal)} lignes de journal, "
          f"{len(rappro_jours)} jours rapprochés -- à publier par 60_publier_donnees.py")


if __name__ == "__main__":
    main()
